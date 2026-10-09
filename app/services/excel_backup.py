"""Backup as a .zip of plain Excel files that anyone can open, which can also be restored.

The zip holds
  00 READ ME.txt                     what each file is
  01 Business details.xlsx ...       one workbook per area, columns in the same format as Utilities → Import
  backup.gstbak                      the exact backup (everything, including payroll, documents, images)
  manifest.json                      fingerprint of every Excel file as written

Restore (always a NEW company — live data is never overwritten):
  * Excel files unchanged → the exact backup inside is restored.
  * Excel files edited, or a zip without the backup → the company is rebuilt from the Excel files through the normal
    entry logic (taxes, stock and accounts recalculated). Files that cannot be imported (attendance, payroll) are
    for reading only. The whole restore is one transaction: any error and nothing is saved.
"""

import datetime as dt
import gzip
import hashlib
import io
import json
import zipfile
from decimal import Decimal

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select, update
from sqlalchemy.orm import Session

from ..deps import Ctx
from ..gst.constants import AccountType, BusinessGstType, Role
from ..gst.states import STATES
from ..models import Account, Business, Employee, Membership, User, Voucher
from ..permissions import effective
from ..schemas import BusinessIn
from . import backup as bk
from . import importer as IM

FORMAT = "gst-billing-excel-backup"
BACKUP_NAME = "backup.gstbak"
HEAD_FILL = PatternFill("solid", fgColor="1F65BB")
INFO_FILL = PatternFill("solid", fgColor="E5E7EB")
D = Decimal

# ---------------------------------------------------------------- columns of the files the importer does not know
BUSINESS_FIELDS = [
    IM.F("name", "Business Name", True), IM.F("legal_name", "Legal Name"), IM.F("gst_type", "GST Registration"),
    IM.F("gstin", "GSTIN"), IM.F("pan", "PAN"), IM.F("state", "State", True), IM.F("address", "Address"),
    IM.F("city", "City"), IM.F("pincode", "Pincode"), IM.F("phone", "Phone"), IM.F("email", "Email"),
    IM.F("entity_type", "Entity Type"), IM.F("gst_registration_date", "GST Registration Date"),
    IM.F("bank_name", "Bank Name"), IM.F("bank_account_no", "Bank Account No"), IM.F("bank_ifsc", "IFSC"),
    IM.F("bank_branch", "Bank Branch"), IM.F("upi_id", "UPI ID"), IM.F("invoice_terms", "Invoice Terms"),
]
ACCOUNT_FIELDS = [
    IM.F("name", "Account Name", True), IM.F("type", "Type", False, "Cash / Bank"), IM.F("bank_name", "Bank Name"),
    IM.F("account_no", "Account No"), IM.F("ifsc", "IFSC"), IM.F("opening_balance", "Opening Balance"),
    IM.F("opening_date", "Opening Date"),
]
STAFF_FIELDS = [
    IM.F("code", "Employee Code"), IM.F("name", "Name", True), IM.F("phone", "Phone"), IM.F("email", "Email"),
    IM.F("designation", "Designation"), IM.F("department", "Department"), IM.F("joined_on", "Joined On"),
    IM.F("left_on", "Left On"), IM.F("salary_type", "Salary Type", False, "Monthly / Daily"),
    IM.F("salary", "Salary"), IM.F("basic_pct", "Basic %"), IM.F("ot_rate", "Overtime Rate"), IM.F("pf", "PF"),
    IM.F("esi", "ESI"), IM.F("pt_monthly", "Professional Tax / Month"), IM.F("tds_monthly", "TDS / Month"),
    IM.F("uan", "UAN"), IM.F("esic_no", "ESIC No"), IM.F("pan", "PAN"), IM.F("bank_name", "Bank Name"),
    IM.F("bank_account", "Bank Account"), IM.F("bank_ifsc", "Bank IFSC"), IM.F("notes", "Notes"),
]
ENTITY_FIELDS = {"business": BUSINESS_FIELDS, "accounts": ACCOUNT_FIELDS, "staff": STAFF_FIELDS,
                 **{k: v[0] for k, v in IM.ENTITIES.items()}}

# file name → (entity, description); restore follows this order (masters before documents before payments)
VOUCHER_FILES = [("purchases", "PURCHASE", "Purchase bills"), ("debit-notes", "PURCHASE_RETURN", "Debit notes (purchase returns)"),
                 ("sales", "SALE", "Sales invoices"), ("credit-notes", "SALE_RETURN", "Credit notes (sales returns)"),
                 ("estimates", "ESTIMATE", "Estimates and quotations"), ("sale-orders", "SALE_ORDER", "Sale orders"),
                 ("delivery-challans", "DELIVERY_CHALLAN", "Delivery challans"),
                 ("purchase-orders", "PURCHASE_ORDER", "Purchase orders"), ("expenses", "EXPENSE", "Expenses")]
READ_ONLY = {"attendance", "payroll"}
EXTRA = "  (for reading)"  # suffix of columns the restore ignores


# ================================================================ who may use it (super admin)
SETTING = "excel_backup"


def platform_default(db: Session) -> bool:
    from ..models import PlatformSetting

    row = db.get(PlatformSetting, SETTING)
    return bool(((row.value or {}) if row else {}).get("default", True))


def set_platform_default(db: Session, on: bool) -> None:
    from ..models import PlatformSetting

    row = db.get(PlatformSetting, SETTING) or PlatformSetting(key=SETTING)
    row.value = {"default": bool(on)}
    db.add(row)


def allowed(db: Session, biz: Business) -> bool:
    return platform_default(db) if biz.excel_backup is None else bool(biz.excel_backup)


def allowed_for_user(db: Session, user: User) -> bool:
    """Restoring a .zip: any business the user owns has Excel backups on (or the platform default, with none)."""
    owned = db.scalars(select(Business).where(Business.owner_id == user.id)).all()
    return any(allowed(db, b) for b in owned) if owned else platform_default(db)


def require(db: Session, biz: Business) -> None:
    if not allowed(db, biz):
        raise HTTPException(403, "Excel backups are not enabled for this business — ask the platform administrator.")


# ================================================================ helpers
def _d(v) -> dt.date | None:
    if not v:
        return None
    return dt.date.fromisoformat(str(v)[:10])


def _n(v) -> float | None:
    return None if v in (None, "") else float(D(str(v)))


def _yn(v) -> str:
    return "Y" if v else "N"


def _state(code) -> str | None:
    return STATES.get(code, code) if code else None


def _title(s: str | None) -> str | None:
    return s.replace("_", " ").title() if s else None


def _workbook(title: str, about: str, fields: list, extra: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    heads = [f.header for f in fields] + [h + EXTRA for h in extra]
    ws.append(heads)
    for i, c in enumerate(ws[1], 1):
        c.font = Font(bold=True, color="FFFFFF" if i <= len(fields) else "374151")
        c.fill = HEAD_FILL if i <= len(fields) else INFO_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = max(12, min(40, len(heads[i - 1]) + 3))
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, dt.date):
                c.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    info = wb.create_sheet("About this file")
    info.append([title])
    info["A1"].font = Font(bold=True, size=13)
    info.append([about])
    info.append([])
    info.append(["Blue columns can be changed and are read back when this backup is restored. Grey columns marked "
                 "'(for reading)' are worked out by the app and are ignored on restore."])
    info.append(["Dates are DD/MM/YYYY. Y = yes, N = no. Do not rename the 'Data' sheet or its column headings."])
    info.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ================================================================ export (from a backup snapshot)
def build_zip(blob: bytes) -> bytes:
    """Excel zip from a company backup (.gstbak bytes)."""
    data = json.loads(gzip.decompress(blob))
    if data.get("format") != bk.FORMAT:
        raise HTTPException(400, "This is not a company backup")
    T = data["tables"]
    biz = (T.get("businesses") or [{}])[0]
    by = lambda name: {r["id"]: r for r in T.get(name, [])}  # noqa: E731
    parties, items, vouchers, accounts = by("parties"), by("items"), by("vouchers"), by("accounts")
    cats, exp_items, employees = by("expense_categories"), by("expense_items"), by("employees")
    lines: dict[str, list] = {}
    for ln in sorted(T.get("voucher_lines", []), key=lambda x: x.get("sort_order") or 0):
        lines.setdefault(ln["voucher_id"], []).append(ln)

    files: list[tuple[str, str, bytes]] = []  # (file name, entity, bytes)

    def add(entity: str, title: str, about: str, rows: list[list], extra: list[str] = (), fields=None, always=False):
        if rows or always:
            files.append((f"{len(files) + 1:02d} {title}.xlsx", entity,
                          _workbook(title, about, fields or ENTITY_FIELDS[entity], list(extra), rows)))

    # ---- business, accounts
    add("business", "Business details", "Your business name, GST registration, address and bank details.", [[
        biz.get("name"), biz.get("legal_name"), _title(biz.get("gst_type")), biz.get("gstin"), biz.get("pan"),
        _state(biz.get("state_code")), biz.get("address"), biz.get("city"), biz.get("pincode"), biz.get("phone"),
        biz.get("email"), _title(biz.get("entity_type")), _d(biz.get("gst_registration_date")), biz.get("bank_name"),
        biz.get("bank_account_no"), biz.get("bank_ifsc"), biz.get("bank_branch"), biz.get("upi_id"), biz.get("invoice_terms")]],
        always=True)
    add("accounts", "Cash and bank accounts", "Cash in hand and bank accounts with their opening balances.",
        [[a["name"], _title(a["type"]), a.get("bank_name"), a.get("account_no"), a.get("ifsc"), _n(a.get("opening_balance")),
          _d(a.get("opening_date"))] for a in accounts.values()])

    # ---- masters
    def party_row(p):
        bal = D(str(p.get("opening_balance") or 0))
        return [p["name"], _title(p["type"]), p.get("gstin"), _title(p.get("gst_type")), p.get("phone"), p.get("email"),
                _state(p.get("state_code")), p.get("billing_address"), p.get("city"), p.get("pincode"),
                p.get("shipping_address"), float(abs(bal)) if bal else None,
                ("Payable" if bal < 0 else "Receivable") if bal else None, _n(p.get("credit_limit"))]
    add("parties", "Customers and suppliers", "Every party with contact details, GST details and opening balance.",
        [party_row(p) for p in parties.values()])

    opening: dict[str, list] = {}
    for m in T.get("stock_movements", []):
        if m["type"] == "OPENING":
            o = opening.setdefault(m["item_id"], [D(0), None])
            o[0] += D(str(m["qty"]))
            o[1] = min(filter(None, [o[1], _d(m["date"])]))
    stock_now: dict[str, D] = {}
    for m in T.get("stock_movements", []):
        stock_now[m["item_id"]] = stock_now.get(m["item_id"], D(0)) + D(str(m["qty"]))
    add("items", "Items and opening stock", "Products and services with prices, GST rate and opening stock.",
        [[i["name"], i.get("code"), _title(i["type"]), i.get("hsn_sac"), i.get("unit"), i.get("category"),
          _n(i.get("sale_price")), _yn(i.get("sale_price_tax_inclusive")), _n(i.get("purchase_price")),
          _yn(i.get("purchase_price_tax_inclusive")), _n(i.get("mrp")), _n(i.get("gst_rate")), _n(i.get("cess_rate")),
          float(opening[i["id"]][0]) if i["id"] in opening else None, opening[i["id"]][1] if i["id"] in opening else None,
          _n(i.get("low_stock_level")), _yn(i.get("track_batch")), _yn(i.get("track_serial")), i.get("description"),
          float(stock_now.get(i["id"], 0)) if i["type"] == "GOODS" else None]
         for i in items.values()], extra=["Stock Now"])
    add("hsn", "HSN and SAC codes", "Your own HSN / SAC list with GST rates.",
        [[h["code"], h.get("description"), _n(h.get("gst_rate")), _n(h.get("cess_rate")), _d(h.get("effective_from"))]
         for h in T.get("hsn_codes", [])])
    add("expense-items", "Expense items", "Expense heads you use, by category.",
        [[e["name"], (cats.get(e.get("category_id")) or {}).get("name"), e.get("hsn_sac"), _n(e.get("rate")), _n(e.get("gst_rate"))]
         for e in exp_items.values()])

    # ---- documents (one row per line)
    for entity, vtype, title in VOUCHER_FILES:
        fields = IM.ENTITIES[entity][0]
        keys = [f.key for f in fields]
        rows = []
        for v in sorted((x for x in vouchers.values() if x["type"] == vtype), key=lambda x: (x["date"], x["number"])):
            party = parties.get(v.get("party_id")) or {}
            orig = vouchers.get(v.get("original_voucher_id")) or {}
            cat = cats.get(v.get("expense_category_id")) or {}
            for ln in lines.get(v["id"], []):
                item = items.get(ln.get("item_id")) or {}
                head = dict(
                    doc_no=v["number"], date=_d(v["date"]), party=party.get("name"), party_gstin=v.get("party_gstin") if party else None,
                    party_phone=v.get("party_phone") if party else None, pos=_state(v.get("place_of_supply")),
                    supplier_bill_no=v.get("supplier_invoice_no"), supplier_bill_date=_d(v.get("supplier_invoice_date")),
                    gst_charged=_yn(v.get("tax_applicable")) if vtype in ("PURCHASE", "PURCHASE_RETURN", "EXPENSE") else None,
                    reverse_charge=_yn(v.get("reverse_charge")) if vtype == "PURCHASE" else None,
                    original_doc=orig.get("number"), reason=v.get("reason"),
                    item=ln["name"], item_code=item.get("code"), hsn_sac=ln.get("hsn_sac"), qty=_n(ln.get("qty")),
                    unit=ln.get("unit"), rate=_n(ln.get("rate")), tax_inclusive=_yn(ln.get("tax_inclusive")),
                    discount_pct=_n(ln.get("discount_pct")), gst_rate=_n(ln.get("gst_rate")), cess_rate=_n(ln.get("cess_rate")),
                    batch_no=ln.get("batch_no"), expiry_date=_d(ln.get("expiry_date")), serial_nos=ln.get("serial_nos"),
                    due_date=_d(v.get("due_date")), notes=v.get("notes"), tcs_rate=_n(v.get("tcs_rate")) or None,
                    category=cat.get("name"), category_type=_title(cat.get("kind")),
                    # payments are in their own files (walk-in bills without a party are marked paid on restore)
                    amount_paid=None, payment_mode=None, account=None)
                tax = sum(D(str(ln.get(k) or 0)) for k in ("cgst", "sgst", "igst", "cess"))
                rows.append([head.get(k) for k in keys] + [
                    _n(ln.get("taxable")), float(tax), _n(ln.get("total")), _n(v.get("grand_total")),
                    "Cancelled" if v.get("cancelled") else "Active", v.get("irn"), v.get("ewb_no")])
        add(entity, title, f"{title}: one row per item line; rows with the same Doc No are one document.", rows,
            extra=["Taxable Value", "Tax", "Line Total", "Document Total", "Status", "IRN", "E-way Bill No"])

    # ---- stock adjustments, payments
    add("stock", "Stock adjustments", "Manual stock corrections (stock from bills is rebuilt from the bill files).",
        [[(items.get(m["item_id"]) or {}).get("name"), (items.get(m["item_id"]) or {}).get("code"), _d(m["date"]),
          _n(m["qty"]), m.get("note")] for m in T.get("stock_movements", []) if m["type"] == "ADJUSTMENT"])
    alloc: dict[str, list] = {}
    for a in T.get("payment_allocations", []):
        alloc.setdefault(a["payment_id"], []).append(a)
    for entity, ptype, title in (("payments-in", "IN", "Payments received"), ("payments-out", "OUT", "Payments made")):
        rows = []
        for p in sorted((x for x in T.get("payments", []) if x["type"] == ptype and x.get("party_id")
                                and x.get("cheque_status") != "BOUNCED"), key=lambda x: x["date"]):  # bounced = void
            party = parties.get(p["party_id"]) or {}
            against = [(vouchers.get(a["voucher_id"]) or {}).get("number") for a in alloc.get(p["id"], [])]
            rows.append([_d(p["date"]), party.get("name"), party.get("gstin"), _n(p["amount"]), _n(p.get("tds_amount")) or None,
                         _title(p.get("mode")), (accounts.get(p.get("account_id")) or {}).get("name"), p.get("reference"),
                         against[0] if against else None, p.get("notes"), p.get("number"), ", ".join(filter(None, against)) or None,
                         _title(p.get("cheque_status"))])
        add(entity, title, f"{title}, with the bills they settle.", rows,
            extra=["Payment No", "Settles Bills", "Cheque Status"])

    # ---- staff
    add("staff", "Staff", "Employees with salary, PF / ESI and bank details.",
        [[e.get("code"), e["name"], e.get("phone"), e.get("email"), e.get("designation"), e.get("department"),
          _d(e.get("joined_on")), _d(e.get("left_on")), _title(e.get("salary_type")), _n(e.get("salary")), _n(e.get("basic_pct")),
          _n(e.get("ot_rate")), _yn(e.get("pf")), _yn(e.get("esi")), _n(e.get("pt_monthly")), _n(e.get("tds_monthly")),
          e.get("uan"), e.get("esic_no"), e.get("pan"), e.get("bank_name"), e.get("bank_account"), e.get("bank_ifsc"), e.get("notes")]
         for e in employees.values()])
    read_fields = lambda *hs: [IM.F(h.lower().replace(" ", "_"), h) for h in hs]  # noqa: E731
    add("attendance", "Attendance (for reading)", "Daily attendance of every employee. Restored only from the exact backup.",
        [[(employees.get(a["employee_id"]) or {}).get("name"), _d(a["date"]), a.get("status"), a.get("check_in"),
          a.get("check_out"), _n(a.get("ot_hours")), a.get("note")]
         for a in sorted(T.get("attendance", []), key=lambda x: (x["date"], x["employee_id"]))],
        fields=read_fields("Employee", "Date", "Status", "Check In", "Check Out", "OT Hours", "Note"))
    runs = by("payroll_runs")
    add("payroll", "Payroll (for reading)", "Salary of every employee by month. Restored only from the exact backup.",
        [[(runs.get(pl["run_id"]) or {}).get("month"), pl.get("name"), _n(pl.get("gross")), _n(pl.get("deductions")),
          _n(pl.get("net")), (runs.get(pl["run_id"]) or {}).get("status")] for pl in T.get("payroll_lines", [])],
        fields=read_fields("Month", "Employee", "Gross", "Deductions", "Net Pay", "Status"))

    readme = [f"Backup of {biz.get('name')} taken on {data.get('created_at', '')[:16].replace('T', ' ')} UTC", "",
              "Each Excel file below holds one part of your business data. Open them in Excel, Google Sheets or LibreOffice.",
              "", *[f"  {name}" for name, _, _ in files], "",
              "backup.gstbak is the exact backup (everything, including documents, images and payroll).", "",
              "To restore: Utilities -> Backup & Restore -> Restore, and choose this .zip file.",
              "  - If you have not changed the Excel files, the exact backup is restored.",
              "  - If you changed them (or made them yourself), the company is rebuilt from the Excel files.",
              "A restore always creates a NEW company; your current data is never overwritten."]
    manifest = {"format": FORMAT, "version": 1, "business_name": biz.get("name"), "created_at": data.get("created_at"),
                "files": {name: {"entity": entity, "sha256": hashlib.sha256(b).hexdigest()} for name, entity, b in files}}
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("00 READ ME.txt", "\r\n".join(readme))
        for name, _, b in files:
            z.writestr(name, b)
        z.writestr(BACKUP_NAME, blob)
        z.writestr("manifest.json", json.dumps(manifest, indent=2))
    return out.getvalue()


def zip_name(business_name: str, created: dt.datetime) -> str:
    return bk.filename(business_name, created).replace(".gstbak", "-excel.zip")


# ================================================================ restore
def is_zip(blob: bytes) -> bool:
    return blob[:2] == b"PK"


def _entity_of(name: str, manifest: dict) -> str | None:
    if name in manifest.get("files", {}):
        return manifest["files"][name]["entity"]
    low = name.lower()
    # most specific first: "Expense items" before "items", "Credit notes (sales returns)" before "sales" …
    guesses = [("attendance", "attendance"), ("payroll", "payroll"), ("business", "business"), ("bank", "accounts"),
               ("expense item", "expense-items"), ("hsn", "hsn"), ("credit note", "credit-notes"), ("debit note", "debit-notes"),
               ("purchase order", "purchase-orders"), ("sale order", "sale-orders"), ("estimate", "estimates"),
               ("quotation", "estimates"), ("challan", "delivery-challans"), ("payments received", "payments-in"),
               ("payments made", "payments-out"), ("stock adjust", "stock"), ("customer", "parties"), ("supplier", "parties"),
               ("part", "parties"), ("item", "items"), ("purchase", "purchases"), ("sales", "sales"), ("invoice", "sales"),
               ("expense", "expenses"), ("staff", "staff"), ("employee", "staff")]
    return next((e for k, e in guesses if k in low), None)


ORDER = ["business", "accounts", "hsn", "expense-items", "parties", "items", *[e for e, _, _ in VOUCHER_FILES],
         "stock", "payments-in", "payments-out", "staff"]


def restore_zip(db: Session, user: User, blob: bytes, name: str | None = None, mode: str = "auto",
                dry_run: bool = False) -> dict:
    try:
        z = zipfile.ZipFile(io.BytesIO(blob))
    except zipfile.BadZipFile as e:
        raise HTTPException(400, "This is not a valid .zip file") from e
    # files may sit inside a folder (zipping the extracted folder); Mac zips add __MACOSX/ copies
    names = [n for n in z.namelist() if not n.endswith("/") and "__MACOSX" not in n]
    base = lambda n: n.replace("\\", "/").rsplit("/", 1)[-1]  # noqa: E731
    man = next((n for n in names if base(n) == "manifest.json"), None)
    try:
        manifest = json.loads(z.read(man)) if man else {}
    except ValueError:
        manifest = {}
    xlsx = {base(n): z.read(n) for n in names if n.lower().endswith(".xlsx") and not base(n).startswith("~$")}
    exact = next((z.read(n) for n in names if n.lower().endswith(".gstbak")), None)
    edited = [n for n, b in xlsx.items()
              if hashlib.sha256(b).hexdigest() != (manifest.get("files", {}).get(n) or {}).get("sha256")]
    if mode == "exact" or (mode == "auto" and exact and not edited):
        if not exact:
            raise HTTPException(400, "This zip has no exact backup inside — restore from the Excel files instead.")
        if dry_run:
            return {"method": "exact", "ok": True, "edited": edited, "errors": []}
        b = bk.restore_as_new(db, exact, user.id, name)
        return {"method": "exact", "ok": True, "business_id": b.id, "business_name": b.name, "edited": edited, "errors": []}
    return _restore_excel(db, user, xlsx, manifest, name, dry_run, edited)


def _rows(entity: str, fname: str, content: bytes) -> list[tuple[int, dict]]:
    fields = ENTITY_FIELDS[entity]
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
    raw = [list(r) for r in ws.iter_rows(values_only=True)]
    if not raw:
        return []
    heads = [str(h or "") for h in raw[0]]
    status_col = next((i for i, h in enumerate(heads) if h.startswith("Status")), None)
    by_header = {IM._norm(f.header): f.key for f in fields}
    cols = [None if h.endswith(EXTRA) else by_header.get(IM._norm(h)) for h in heads]
    if not any(cols):
        raise HTTPException(400, f"{fname}: column headings do not match — keep the headings as they were exported")
    out = []
    for idx, r in enumerate(raw[1:], start=2):
        if status_col is not None and status_col < len(r) and str(r[status_col] or "").lower() == "cancelled":
            continue  # cancelled documents are not re-created
        rec = {k: (v.strip() if isinstance(v, str) else v) for k, v in zip(cols, r) if k and v not in (None, "")}
        rec = {k: v for k, v in rec.items() if v != ""}
        if rec:
            out.append((idx, rec))
    return out


def _restore_excel(db: Session, user: User, xlsx: dict[str, bytes], manifest: dict, name: str | None, dry_run: bool,
                   edited: list[str]) -> dict:
    plan: dict[str, list[tuple[str, bytes]]] = {}
    skipped = []
    for fname, content in xlsx.items():
        ent = _entity_of(fname, manifest)
        if ent in ENTITY_FIELDS:
            plan.setdefault(ent, []).append((fname, content))
        else:
            skipped.append(fname)
    if "business" not in plan:
        raise HTTPException(400, "The zip needs a 'Business details' Excel file to rebuild the company.")
    errors: list[dict] = []
    summary: dict[str, str] = {}
    try:
        # ---- the company
        (bname, bcontent), = plan["business"][:1]
        rows = _rows("business", bname, bcontent)
        if not rows:
            raise HTTPException(400, "Business details: the first data row is empty")
        r = rows[0][1]
        try:
            data = BusinessIn(name=name or f"{r.get('name')} (restored)", legal_name=r.get("legal_name"),
                              gst_type=IM.enum_of(BusinessGstType, r.get("gst_type"), BusinessGstType.UNREGISTERED),
                              gstin=str(r["gstin"]).upper() if r.get("gstin") else None, pan=r.get("pan"),
                              state_code=IM.state_code(r.get("state")) or "27", address=r.get("address"), city=r.get("city"),
                              pincode=str(r["pincode"]) if r.get("pincode") else None, phone=str(r["phone"]) if r.get("phone") else None,
                              email=r.get("email"), bank_name=r.get("bank_name"),
                              bank_account_no=str(r["bank_account_no"]) if r.get("bank_account_no") else None,
                              bank_ifsc=r.get("bank_ifsc"), bank_branch=r.get("bank_branch"), upi_id=r.get("upi_id"),
                              invoice_terms=r.get("invoice_terms"),
                              entity_type=str(r.get("entity_type") or "Proprietorship").upper().replace(" ", "_"),
                              gst_registration_date=IM.date(r.get("gst_registration_date"), "GST Registration Date"))
        except (IM.RowError, ValueError) as e:
            raise HTTPException(400, f"{bname}: {IM._pyd_err(e) if hasattr(e, 'errors') else e}") from e
        from ..routers.businesses import apply_business
        from ..routers.expenses import seed_categories
        from .accounts import cash_account
        from .godowns import default_godown
        from .plans import start_trial

        start_trial(db, user.id)
        biz = Business(owner_id=user.id)
        apply_business(biz, data)
        biz.stock_control = "ALLOW"  # historical documents are re-created as they were
        db.add(biz)
        db.flush()
        m = Membership(user_id=user.id, business_id=biz.id, role=Role.OWNER)
        db.add(m)
        cash = cash_account(db, biz.id)
        seed_categories(db, biz.id)
        default_godown(db, biz.id)
        db.flush()
        ctx = Ctx(db=db, user=user, business=biz, role=Role.OWNER, membership=m, perms=effective(Role.OWNER, None),
                  restoring=True)
        summary["business"] = biz.name

        for ent in ORDER[1:]:
            for fname, content in plan.get(ent, []):
                try:
                    rows = _rows(ent, fname, content)
                except HTTPException as e:
                    errors.append({"file": fname, "row": None, "message": e.detail})
                    continue
                if not rows:
                    continue
                if ent == "accounts":
                    errs = _import_accounts(ctx, rows, cash)
                    summary[fname] = f"{len(rows)} accounts"
                elif ent == "staff":
                    errs = _import_staff(ctx, rows)
                    summary[fname] = f"{len(rows)} employees"
                else:
                    try:
                        s, _, errs = IM.import_rows(ctx, ent, rows)
                    except Exception as e:  # noqa: BLE001 — never a server error for a bad file
                        s, errs = "", [{"row": None, "message": f"Could not read this file ({type(e).__name__}: {e})"}]
                    summary[fname] = s
                errors += [{"file": fname, **e} for e in errs]
        if not errors:
            # restored documents count in the month they were made, not today (plan limits)
            db.execute(update(Voucher).where(Voucher.business_id == biz.id).values(created_at=Voucher.date))
            db.flush()
    except HTTPException:
        db.rollback()
        raise
    ok = not errors
    if ok and not dry_run:
        biz.stock_control = "WARN"
        db.commit()
        return {"method": "excel", "ok": True, "business_id": biz.id, "business_name": biz.name, "summary": summary,
                "edited": edited, "skipped": skipped, "errors": []}
    db.rollback()
    return {"method": "excel", "ok": ok, "dry_run": dry_run, "summary": summary, "edited": edited, "skipped": skipped,
            "errors": errors[:500]}


def _import_accounts(ctx: Ctx, rows, cash: Account) -> list[dict]:
    errs = []
    for n, r in rows:
        try:
            atype = IM.enum_of(AccountType, r.get("type"), AccountType.BANK)
            bal = IM.num(r.get("opening_balance"), "Opening Balance", D(0))
            when = IM.date(r.get("opening_date"), "Opening Date")
            existing = ctx.db.scalar(select(Account).where(Account.business_id == ctx.bid, Account.name == str(r["name"])))
            if existing is None and atype == AccountType.CASH and cash.name.lower() == str(r["name"]).lower():
                existing = cash
            if existing is None:
                existing = Account(business_id=ctx.bid, type=atype, name=str(r["name"]))
                ctx.db.add(existing)
            existing.bank_name, existing.ifsc = r.get("bank_name"), r.get("ifsc")
            existing.account_no = str(r["account_no"]) if r.get("account_no") else None
            existing.opening_balance = bal
            if when:
                existing.opening_date = when
            ctx.db.flush()
        except IM.RowError as e:
            errs.append({"row": n, "message": str(e)})
    return errs


def _import_staff(ctx: Ctx, rows) -> list[dict]:
    errs = []
    for n, r in rows:
        try:
            st = str(r.get("salary_type") or "Monthly").upper()
            if st not in ("MONTHLY", "DAILY"):
                raise IM.RowError("Salary Type: use Monthly or Daily")
            ctx.db.add(Employee(
                business_id=ctx.bid, code=str(r["code"]) if r.get("code") else None, name=str(r["name"]),
                phone=str(r["phone"]) if r.get("phone") else None, email=r.get("email"), designation=r.get("designation"),
                department=r.get("department"), joined_on=IM.date(r.get("joined_on"), "Joined On"),
                left_on=IM.date(r.get("left_on"), "Left On"), is_active=not r.get("left_on"), salary_type=st,
                salary=IM.num(r.get("salary"), "Salary", D(0)), basic_pct=IM.num(r.get("basic_pct"), "Basic %", D(50)),
                ot_rate=IM.num(r.get("ot_rate"), "Overtime Rate", D(0)), pf=bool(IM.yes(r.get("pf"))), esi=bool(IM.yes(r.get("esi"))),
                pt_monthly=IM.num(r.get("pt_monthly"), "Professional Tax", D(0)), tds_monthly=IM.num(r.get("tds_monthly"), "TDS", D(0)),
                uan=str(r["uan"]) if r.get("uan") else None, esic_no=str(r["esic_no"]) if r.get("esic_no") else None,
                pan=r.get("pan"), bank_name=r.get("bank_name"), bank_account=str(r["bank_account"]) if r.get("bank_account") else None,
                bank_ifsc=r.get("bank_ifsc"), notes=r.get("notes")))
            ctx.db.flush()
        except IM.RowError as e:
            errs.append({"row": n, "message": str(e)})
    return errs
