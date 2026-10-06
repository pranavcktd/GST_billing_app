"""E-way bills without an API: which bills need one, bulk JSON for the portal, importing the generated
numbers back, and validity tracking.

Validity (Rule 138(10)): one day for every 200 km (20 km for over-dimensional cargo) or part of it; a
day ends at midnight. The km per day is configurable (Admin → GST config).
"""

import csv
import datetime as dt
import io
import math
import re

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..gst.constants import VoucherType
from ..models import Business, Voucher
from . import config_store
from .einvoice import EWB_DOC_TYPES, IST, PayloadError, ewaybill_payload

TYPES = list(EWB_DOC_TYPES)


def valid_till(generated: dt.datetime, distance_km: float | int | None, odc: bool = False) -> dt.datetime | None:
    if not distance_km:
        return None
    per_day = 20 if odc else int(config_store.get("ewb_km_per_day"))
    days = max(1, math.ceil(float(distance_km) / per_day))
    local = generated.astimezone(IST) if generated.tzinfo else generated.replace(tzinfo=IST)
    return dt.datetime.combine(local.date() + dt.timedelta(days=days), dt.time(23, 59, 59), tzinfo=IST)


def is_goods(v: Voucher) -> bool:
    return any(not (l.hsn_sac or "").startswith("99") for l in v.lines)


def needs_ewb(v: Voucher) -> bool:
    """Goods moving with value above the threshold and no e-way bill recorded yet."""
    if v.cancelled or v.type not in TYPES or v.ewb_no or not is_goods(v):
        return False
    return v.grand_total > config_store.get("ewb_threshold", v.date)


def status(v: Voucher, now: dt.datetime) -> str:
    if v.ewb_no:
        if not v.ewb_valid_till:
            return "ACTIVE"
        till = v.ewb_valid_till if v.ewb_valid_till.tzinfo else v.ewb_valid_till.replace(tzinfo=IST)
        if till < now:
            return "EXPIRED"
        return "EXPIRING" if till - now < dt.timedelta(hours=36) else "ACTIVE"
    return "PENDING" if needs_ewb(v) else "NOT_NEEDED"


def listing(db: Session, biz: Business, date_from: dt.date, date_to: dt.date) -> list[dict]:
    now = dt.datetime.now(IST)
    out = []
    for v in db.scalars(select(Voucher).where(Voucher.business_id == biz.id, Voucher.type.in_(TYPES), Voucher.cancelled.is_(False),
                                              Voucher.date >= date_from, Voucher.date <= date_to).order_by(Voucher.date.desc())):
        st = status(v, now)
        if st == "NOT_NEEDED":
            continue
        problems = []
        if st == "PENDING":
            try:
                ewaybill_payload(db, biz, v)
            except PayloadError as e:
                problems = e.problems
        t = v.transport or {}
        out.append(dict(id=v.id, type=v.type.value, number=v.number, date=v.date, party=v.party_name, value=float(v.grand_total),
                        status=st, ewb_no=v.ewb_no, ewb_date=v.ewb_date, valid_till=v.ewb_valid_till,
                        vehicle=t.get("vehicle_no"), distance=t.get("distance_km"), ready=st == "PENDING" and not problems,
                        problems=problems))
    return out


def bulk(db: Session, biz: Business, ids: list[str]) -> tuple[dict, list[dict]]:
    bills, skipped = [], []
    for v in db.scalars(select(Voucher).where(Voucher.business_id == biz.id, Voucher.id.in_(ids))):
        try:
            bills += ewaybill_payload(db, biz, v)["billLists"]
        except PayloadError as e:
            skipped.append(dict(number=v.number, problems=e.problems))
    return {"version": "1.0.0621", "billLists": bills}, skipped


# ---------------------------------------------------------------- import generated numbers from the portal
EWB_COLS = {"ewb no", "ewb no.", "ewb number", "e-way bill no", "e-way bill no.", "eway bill no", "e way bill no", "ewbno", "ewaybillno"}
DATE_COLS = {"ewb date", "ewb generated date", "generated date", "generation date", "ewaybilldate", "ewb date & time", "date"}
DOC_COLS = {"doc no", "doc no.", "document no", "document no.", "doc number", "document number", "docno", "invoice no", "bill no"}
VALID_COLS = {"valid upto", "valid up to", "valid until", "validity", "valid till", "validupto", "valid upto date"}


def _n(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def _dt(v) -> dt.datetime | None:
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v if v.tzinfo else v.replace(tzinfo=IST)
    if isinstance(v, dt.date):
        return dt.datetime.combine(v, dt.time(0, 0), tzinfo=IST)
    s = str(v).strip()
    for fmt in ("%d/%m/%Y %I:%M:%S %p", "%d/%m/%Y %I:%M %p", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
                "%d-%m-%Y %H:%M:%S", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s, fmt).replace(tzinfo=IST)
        except ValueError:
            pass
    return None


def parse_import(filename: str, content: bytes) -> list[dict]:
    if filename.lower().endswith(".csv"):
        table = list(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))
    else:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        table = [list(r) for r in wb.active.iter_rows(values_only=True)]
    for i, row in enumerate(table[:30]):
        names = [_n(c) for c in row]
        cols = {}
        for key, opts in (("ewb", EWB_COLS), ("date", DATE_COLS), ("doc", DOC_COLS), ("valid", VALID_COLS)):
            for j, h in enumerate(names):
                if h in opts and key not in cols:
                    cols[key] = j
        if "ewb" in cols and "doc" in cols:
            out = []
            for r in table[i + 1:]:
                get = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None  # noqa: E731
                ewb = re.sub(r"\D", "", str(get("ewb") or ""))
                doc = str(get("doc") or "").strip()
                if len(ewb) == 12 and doc:
                    out.append(dict(ewb_no=ewb, doc_no=doc, ewb_date=_dt(get("date")), valid_till=_dt(get("valid"))))
            return out
    raise ValueError("Could not find the 'EWB No' and 'Doc No' columns — upload the e-way bill list downloaded from the portal")


def apply_import(db: Session, biz: Business, rows: list[dict]) -> dict:
    vouchers = db.scalars(select(Voucher).where(Voucher.business_id == biz.id, Voucher.type.in_(TYPES), Voucher.cancelled.is_(False))).all()
    by_doc: dict[str, Voucher] = {}
    for v in vouchers:
        key = v.supplier_invoice_no if v.type == VoucherType.PURCHASE and v.supplier_invoice_no else v.number
        by_doc[key.strip().upper()] = v
    updated, unmatched, unchanged = [], [], 0
    for r in rows:
        v = by_doc.get(r["doc_no"].upper())
        if v is None:
            unmatched.append(r["doc_no"])
            continue
        if v.ewb_no == r["ewb_no"] and v.ewb_valid_till:
            unchanged += 1
            continue
        v.ewb_no = r["ewb_no"]
        v.ewb_date = r["ewb_date"] or v.ewb_date or dt.datetime.now(IST)
        v.ewb_valid_till = r["valid_till"] or valid_till(v.ewb_date, (v.transport or {}).get("distance_km"))
        updated.append(dict(number=v.number, ewb_no=v.ewb_no))
    db.commit()
    return {"updated": updated, "unmatched": unmatched, "unchanged": unchanged}
