"""Platform HSN/SAC master: reading official files, the upload template, and code checks.

Accepted uploads
  * Excel (.xlsx) — every sheet is read; the header row is found automatically. Works with the
    GST portal download (sheets HSN_MSTR / SAC_MSTR, columns HSN_CD / SAC_CD + description) and
    with our template (code, description, GST %, cess %, effective from).
  * CSV — same columns.
  * PDF — text is read line by line: a line starting with a 2-8 digit code (optionally after a
    serial number and a label such as "Heading" / "Service Code") starts an entry; following
    lines without a code continue its description. Rates are NEVER read from a PDF (official
    PDFs mix in customs duty and other percentages) — set them with bulk rates or rate notices.

HSN codes have 2/4/6/8 digits; SAC codes start with 99 and have 2/4/5/6 digits.
"""

import csv
import datetime as dt
import io
import re
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..gst.constants import ItemType
from ..models import MasterHsn
from . import config_store

MAX_ROWS = 60000

ALIASES = {
    "code": {"hsn_cd", "sac_cd", "hsn", "sac", "hsn code", "sac code", "hsn/sac code", "hsn/sac", "hsn_sac", "code",
             "service code", "service code (tariff)", "tariff item", "hsn_code", "sac_code", "chapter/heading/sub-heading/tariff item"},
    "description": {"hsn_description", "sac_description", "description", "desc", "service description",
                    "description of goods", "description of services", "description of goods/services", "hsn_desc", "sac_desc"},
    "gst_rate": {"gst %", "gst rate", "gst rate (%)", "rate", "rate (%)", "gst", "igst", "igst %", "igst rate", "igst rate (%)", "tax rate"},
    "cess_rate": {"cess %", "cess", "cess rate", "compensation cess", "cess rate (%)"},
    "effective_from": {"effective from", "effective_from", "effective date", "w.e.f.", "wef", "with effect from"},
}
LABELS = {"chapter", "section", "heading", "sub-heading", "subheading", "group", "service", "code", "tariff", "item",
          "sac", "hsn", "(tariff)", "no.", "no"}


def kind_of(code: str) -> str:
    return "SAC" if code.startswith("99") else "HSN"


def clean_code(value, numeric_cell: bool = False) -> str | None:
    """'0101 21 00', '0101.21.00', 101 (Excel number) → '01012100' / '0101'. None when not a code."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = re.sub(r"[\s.\-]", "", str(value).strip())
    if not s.isdigit():
        return None
    if numeric_cell and len(s) % 2 == 1 and not s.startswith("99"):
        s = "0" + s  # Excel dropped the leading zero of chapters 01-09
    return s if 2 <= len(s) <= 8 else None


def _num(v) -> Decimal | None:
    if v is None or str(v).strip() in ("", "-", "NIL", "Nil", "nil"):
        return None
    try:
        return Decimal(str(v).replace("%", "").strip())
    except InvalidOperation:
        raise ValueError(f"'{v}' is not a number") from None


def _date(v) -> dt.date | None:
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"'{v}' is not a date (use YYYY-MM-DD or DD/MM/YYYY)")


def _header_map(row) -> dict[str, int] | None:
    found = {}
    for i, cell in enumerate(row):
        h = re.sub(r"\s+", " ", str(cell or "").strip().lower())
        for key, names in ALIASES.items():
            if h in names and key not in found:
                found[key] = i
    return found if "code" in found else None


def _table_rows(table, numeric_codes: bool, errors: list, sheet: str):
    """Yield parsed rows from a 2-D table (list of rows)."""
    header = None
    for n, row in enumerate(table, 1):
        if header is None:
            header = _header_map(row)
            continue
        if not any(c not in (None, "") for c in row):
            continue
        cell = lambda k: row[header[k]] if k in header and header[k] < len(row) else None  # noqa: E731
        raw = cell("code")
        code = clean_code(raw, numeric_cell=numeric_codes and isinstance(raw, (int, float)))
        if code is None:
            if raw not in (None, ""):
                errors.append({"row": f"{sheet} {n}".strip(), "error": f"'{raw}' is not a 2-8 digit HSN/SAC code"})
            continue
        try:
            yield dict(code=code, description=(str(cell("description")).strip() if cell("description") not in (None, "") else None),
                       gst_rate=_num(cell("gst_rate")), cess_rate=_num(cell("cess_rate")),
                       effective_from=_date(cell("effective_from")))
        except ValueError as e:
            errors.append({"row": f"{sheet} {n}".strip(), "error": f"{code}: {e}"})
    if header is None:
        errors.append({"row": sheet or "file", "error": "No header row with an HSN/SAC code column was found"})


def _pdf_rows(content: bytes, errors: list):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    current = None
    for page in reader.pages:
        for line in (page.extract_text() or "").splitlines():
            line = re.sub(r"\s+", " ", line).strip()
            if not line:
                continue
            tokens = line.split(" ")
            i = 0
            if tokens and re.fullmatch(r"\d{1,3}[.)]?", tokens[0]) and len(tokens) > 1:
                # a serial number — only skip it when a code or label follows
                nxt = tokens[1].lower()
                if nxt in LABELS or re.fullmatch(r"\d{2,8}", tokens[1]):
                    i = 1
            while i < len(tokens) and tokens[i].lower().strip(":") in LABELS:
                i += 1
            code_parts = []
            while i < len(tokens) and re.fullmatch(r"\d{2,8}", tokens[i]) and len("".join(code_parts) + tokens[i]) <= 8:
                code_parts.append(tokens[i])
                i += 1
                if len(code_parts[0]) != 4:  # only tariff-style '0101 21 00' is split into groups
                    break
            code = clean_code("".join(code_parts)) if code_parts else None
            desc = " ".join(tokens[i:]).lstrip("-–— :").strip()
            if code and desc and not re.fullmatch(r"[\d\s.%]+", desc):
                if current:
                    yield current
                current = dict(code=code, description=desc, gst_rate=None, cess_rate=None, effective_from=None)
            elif current and not code and len(line) > 2 and not re.fullmatch(r"(page )?\d+( of \d+)?", line.lower()):
                if len(current["description"]) < 1500:
                    current["description"] += " " + line
    if current:
        yield current
    if not reader.pages:
        errors.append({"row": "file", "error": "The PDF has no pages"})


def parse_file(filename: str, content: bytes) -> tuple[list[dict], list[dict], str]:
    """Rows (deduplicated by code, last one wins), errors, detected format."""
    name = filename.lower()
    errors: list[dict] = []
    rows: list[dict] = []
    if name.endswith(".pdf"):
        fmt = "PDF"
        rows = list(_pdf_rows(content, errors))
        if not rows:
            errors.append({"row": "file", "error": "No HSN/SAC codes could be read — is the PDF scanned (an image)? "
                                                   "Use the Excel download from the GST portal instead."})
    elif name.endswith(".csv"):
        fmt = "CSV"
        table = list(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))
        rows = list(_table_rows(table, False, errors, ""))
    elif name.endswith((".xlsx", ".xlsm")):
        fmt = "Excel"
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for ws in wb.worksheets:
            if ws.title.lower().startswith(("instruction", "read me", "readme")):
                continue
            table = [list(r) for r in ws.iter_rows(values_only=True)]
            before = len(errors)
            got = list(_table_rows(table, True, errors, ws.title))
            if not got and len(errors) == before + 1 and "No header" in errors[-1]["error"] and len(wb.worksheets) > 1:
                errors.pop()  # an unrelated sheet in a multi-sheet workbook
            rows += got
    elif name.endswith(".xls"):
        raise ValueError("Old .xls files are not supported — open it in Excel and 'Save As' .xlsx")
    else:
        raise ValueError("Upload an Excel (.xlsx), CSV or PDF file")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"The file has more than {MAX_ROWS} codes")
    by_code = {}
    for r in rows:
        by_code[r["code"]] = r
    return list(by_code.values()), errors, fmt


def validate_rates(rows: list[dict], errors: list[dict]) -> list[dict]:
    allowed = config_store.all_rates()
    ok = []
    for r in rows:
        if r["gst_rate"] is not None and r["gst_rate"] not in allowed:
            errors.append({"row": r["code"], "error": f"GST {r['gst_rate']}% is not a configured slab"})
            continue
        ok.append(r)
    return ok


def apply_rows(db: Session, rows: list[dict], dry_run: bool) -> dict:
    """Insert / update codes. A blank rate or description in the file never erases a stored one."""
    existing = {h.code: h for h in db.scalars(select(MasterHsn).where(MasterHsn.code.in_([r["code"] for r in rows])))} if rows else {}
    created = updated = unchanged = rate_changes = 0
    for r in rows:
        h = existing.get(r["code"])
        if h is None:
            created += 1
            if not dry_run:
                db.add(MasterHsn(code=r["code"], description=r["description"], gst_rate=r["gst_rate"],
                                 cess_rate=r["cess_rate"] or Decimal("0"), effective_from=r["effective_from"]))
            continue
        new = dict(description=r["description"] or h.description,
                   gst_rate=r["gst_rate"] if r["gst_rate"] is not None else h.gst_rate,
                   cess_rate=r["cess_rate"] if r["cess_rate"] is not None else h.cess_rate,
                   effective_from=r["effective_from"] or h.effective_from)
        if all(getattr(h, k) == v for k, v in new.items()):
            unchanged += 1
            continue
        updated += 1
        rate_changes += int(new["gst_rate"] != h.gst_rate)
        if not dry_run:
            for k, v in new.items():
                setattr(h, k, v)
    if not dry_run:
        db.flush()
    return dict(created=created, updated=updated, unchanged=unchanged, rate_changes=rate_changes)


def stats(db: Session) -> dict:
    total = db.scalar(select(func.count()).select_from(MasterHsn)) or 0
    sac = db.scalar(select(func.count()).select_from(MasterHsn).where(MasterHsn.code.like("99%"))) or 0
    no_rate = db.scalar(select(func.count()).select_from(MasterHsn).where(MasterHsn.gst_rate.is_(None))) or 0
    return dict(total=total, hsn=total - sac, sac=sac, without_rate=no_rate)


def template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "HSN SAC master"
    head = ["HSN/SAC Code", "Description", "GST %", "Cess %", "Effective From"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F65BB")
    examples = [
        ("61", "Articles of apparel and clothing accessories, knitted or crocheted", None, None, None),
        ("6109", "T-shirts, singlets and other vests, knitted or crocheted", 5, 0, "2025-09-22"),
        ("84713010", "Personal computer (laptop)", 18, 0, "2025-09-22"),
        ("2202", "Waters, including mineral and aerated waters, containing added sugar", 40, 0, "2025-09-22"),
        ("9954", "Construction services", None, None, None),
        ("998719", "Maintenance and repair services of other goods", 18, 0, None),
    ]
    for r in examples:
        ws.append(list(r))
    for col, width in zip("ABCDE", (16, 70, 8, 8, 16)):
        ws.column_dimensions[col].width = width
    info = wb.create_sheet("Instructions")
    for line in [
        "HSN / SAC master — upload template",
        "",
        "HSN/SAC Code (required): 2, 4, 6 or 8 digits for goods (HSN); SAC codes start with 99 (2, 4, 5 or 6 digits).",
        "   Keep the column formatted as Text so leading zeros (e.g. 0101) are not lost.",
        "Description: the official description.",
        "GST %: optional suggested rate. It must be one of the slabs configured in Admin → GST config.",
        "   Leave blank when a code has no single rate (e.g. a chapter heading) — it can be set later.",
        "Cess %: optional compensation cess.",
        "Effective From: optional date (YYYY-MM-DD or DD/MM/YYYY) of the notification the rate comes from.",
        "",
        "Re-uploading updates existing codes. Blank cells never erase a value already stored.",
        "The GST portal's HSN/SAC Excel (sheets HSN_MSTR and SAC_MSTR) can be uploaded as it is.",
        "PDFs are read for codes and descriptions only — rates are never taken from a PDF.",
        "Always use Preview first to check what will change.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 120
    info["A1"].font = Font(bold=True, size=13)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export(db: Session) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "HSN SAC master"
    ws.append(["HSN/SAC Code", "Description", "GST %", "Cess %", "Effective From"])
    for h in db.scalars(select(MasterHsn).order_by(MasterHsn.code)):
        ws.append([h.code, h.description, float(h.gst_rate) if h.gst_rate is not None else None, float(h.cess_rate or 0),
                   h.effective_from.isoformat() if h.effective_from else None])
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ================================================================ checks on items and bills
class HsnError(ValueError):
    pass


def check_code(db: Session, code: str | None, item_type: ItemType | None = None, on: dt.date | None = None) -> None:
    """Raise HsnError for a wrong code: goods/service mismatch, or (strict mode) not in the master."""
    if not code:
        return
    if item_type == ItemType.SERVICE and not code.startswith("99"):
        raise HsnError(f"{code} is an HSN (goods) code — services use SAC codes starting with 99")
    if item_type == ItemType.GOODS and code.startswith("99"):
        raise HsnError(f"{code} is a SAC (services) code — choose item type 'Service' or use an HSN code")
    if config_store.get("hsn_strict", on) and db.get(MasterHsn, code) is None \
            and (db.scalar(select(func.count()).select_from(MasterHsn)) or 0) > 0:
        raise HsnError(f"HSN/SAC {code} is not in the official master. Pick a listed code or request it to be added "
                       "(Utilities → HSN/SAC).")


def lookup(db: Session, code: str) -> dict:
    """Exact match plus its chapter / heading descriptions, for the item form."""
    code = re.sub(r"\D", "", code or "")
    hit = db.get(MasterHsn, code) if code else None
    parents = []
    for n in (2, 4, 5, 6):
        if len(code) > n and (p := db.get(MasterHsn, code[:n])) is not None:
            parents.append(dict(code=p.code, description=p.description))
    return dict(code=code, kind=kind_of(code) if code else None, found=hit is not None,
                description=hit.description if hit else None,
                gst_rate=float(hit.gst_rate) if hit and hit.gst_rate is not None else None,
                cess_rate=float(hit.cess_rate) if hit and hit.cess_rate is not None else None,
                parents=parents, strict=bool(config_store.get("hsn_strict")),
                min_digits=int(config_store.get("hsn_digits_small")))
