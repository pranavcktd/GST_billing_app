"""Getting figures into a practice file: trial-balance files (Excel / CSV, incl. Tally exports) and the
books of a business kept in MyBillSync."""

import csv
import datetime as dt
import io
import re
import uuid
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Business, CapitalEntry
from ..reports.accounting import balance_sheet, profit_and_loss
from .final_accounts import GROUPS, ZERO, D, auto_head, tb_template_rows

NAME_COLS = {"ledger", "ledger name", "particulars", "name", "account", "account name", "head"}
GROUP_COLS = {"group", "under", "parent", "group name"}
DR_COLS = {"debit", "dr", "closing debit", "debit amount", "closing balance debit"}
CR_COLS = {"credit", "cr", "closing credit", "credit amount", "closing balance credit"}
AMT_COLS = {"amount", "closing balance", "balance", "closing"}
PY_DR = {"previous year debit", "py debit", "py dr", "previous debit", "last year debit"}
PY_CR = {"previous year credit", "py credit", "py cr", "previous credit", "last year credit"}
PY_AMT = {"previous year", "previous year amount", "py amount", "last year", "py"}


def new_id() -> str:
    return uuid.uuid4().hex[:12]


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower().replace("(", "").replace(")", "").replace(".", ""))


def _signed(v) -> Decimal:
    """'12,500.00 Dr' → 12500, '8,000 Cr' → -8000, '(500)' → -500."""
    if v is None or v == "":
        return ZERO
    if isinstance(v, (int, float, Decimal)):
        return D(v)
    s = str(v).strip().replace(",", "")
    neg = s.lower().endswith("cr") or (s.startswith("(") and s.endswith(")")) or s.startswith("-")
    s = re.sub(r"(?i)dr|cr|[()₹\s-]", "", s)
    return -D(s) if neg else D(s)


def _table(filename: str, content: bytes) -> list[list]:
    name = filename.lower()
    if name.endswith(".csv"):
        return list(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if any(any(c not in (None, "") for c in r) for r in rows):
                return rows
        return []
    raise ValueError("Upload the trial balance as Excel (.xlsx) or CSV")


def parse_trial_balance(filename: str, content: bytes) -> tuple[list[dict], list[str]]:
    """Ledgers with signed closing balances (debit +, credit −). Group header rows of Tally exports
    (a group name followed by its ledgers adding up to it) are recognised and skipped."""
    table = _table(filename, content)
    header_at, cols = None, {}
    for i, row in enumerate(table[:30]):
        names = [_norm(c) for c in row]
        c = {}
        for j, h in enumerate(names):
            for key, options in (("name", NAME_COLS), ("group", GROUP_COLS), ("dr", DR_COLS), ("cr", CR_COLS),
                                 ("amt", AMT_COLS), ("py_dr", PY_DR), ("py_cr", PY_CR), ("py_amt", PY_AMT)):
                if h in options and key not in c:
                    c[key] = j
        if "name" in c and ({"dr", "cr"} & set(c) or "amt" in c):
            header_at, cols = i, c
            break
    if header_at is None:
        raise ValueError("Could not find the header row — the file needs a Ledger / Particulars column and "
                         "Debit / Credit (or Amount with Dr / Cr) columns. Download the template to see the layout.")
    warnings: list[str] = []
    raw = []
    for row in table[header_at + 1:]:
        get = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None  # noqa: E731
        name = str(get("name") or "").strip()
        if not name or _norm(name) in ("total", "grand total", "difference in opening balances"):
            continue
        if "dr" in cols or "cr" in cols:
            cy = _signed(get("dr")) - _signed(get("cr")).copy_abs() if get("cr") not in (None, "") else _signed(get("dr"))
        else:
            cy = _signed(get("amt"))
        if "py_dr" in cols or "py_cr" in cols:
            py = _signed(get("py_dr")) - (_signed(get("py_cr")).copy_abs() if get("py_cr") not in (None, "") else ZERO)
        else:
            py = _signed(get("py_amt"))
        raw.append(dict(name=name, group=str(get("group") or "").strip() or None, cy=cy, py=py))

    # Tally-style group rows: "Sundry Debtors 90,000" followed by its ledgers adding up to 90,000
    out, i = [], 0
    while i < len(raw):
        r = raw[i]
        if not r["group"] and _norm(r["name"]) in GROUPS:
            j, total = i + 1, ZERO
            while j < len(raw) and _norm(raw[j]["name"]) not in GROUPS and total != r["cy"]:
                total += raw[j]["cy"]
                j += 1
            if j > i + 1 and total == r["cy"]:
                for child in raw[i + 1:j]:
                    child["group"] = r["name"]
                    out.append(child)
                i = j
                continue
        out.append(r)
        i += 1
    ledgers = [dict(id=new_id(), name=r["name"], group=r["group"], head=auto_head(r["name"], r["group"]), partner=None,
                    cy=str(r["cy"]), py=str(r["py"])) for r in out if r["cy"] or r["py"]]
    diff = sum((D(l["cy"]) for l in ledgers), ZERO)
    if diff:
        warnings.append(f"The imported trial balance does not tally (difference {diff}).")
    if not ledgers:
        warnings.append("No ledger balances were found in the file.")
    return ledgers, warnings


def template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Trial balance"
    for r in tb_template_rows():
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F65BB")
    for col, w in zip("ABCDEF", (32, 22, 14, 14, 20, 20)):
        ws.column_dimensions[col].width = w
    info = wb.create_sheet("Instructions")
    for line in ["Trial balance import", "",
                 "One row per ledger with its closing balance for the year: Debit or Credit.",
                 "Group (optional): the Tally / books group, e.g. Sundry Debtors, Indirect Expenses — used to classify ledgers.",
                 "Previous-year columns are optional and give the comparative figures.",
                 "Do not include closing stock — enter it on the Figures tab.",
                 "A Tally trial balance exported to Excel (Particulars / Debit / Credit) can be uploaded as it is."]:
        info.append([line])
    info.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- MyBillSync books → ledgers
def _books_year(db: Session, biz: Business, start: dt.date, end: dt.date) -> tuple[dict[str, tuple[str, Decimal]], Decimal]:
    pl = profit_and_loss(db, biz, start, end)
    bs = balance_sheet(db, biz, end)
    lines: dict[str, tuple[str, Decimal]] = {}

    def put(name, head, amount):
        amount = D(amount)
        if amount:
            prev = lines.get(name, (head, ZERO))[1]
            lines[name] = (head, prev + amount)

    put("Sales", "REVENUE", -pl["sales"])
    put("Sales returns", "REVENUE", pl["sale_returns"])
    put("Opening stock", "OPENING_STOCK", pl["opening_stock"])
    put("Purchases", "PURCHASES", pl["purchases"])
    put("Purchase returns", "PURCHASES", -pl["purchase_returns"])
    for name, amt in pl["direct"].items():
        put(name, "DIRECT_EXP", amt)
    for name, amt in pl["indirect"].items():
        put(name, auto_head(name) if auto_head(name) in ("EMPLOYEE", "FINANCE") else "OTHER_EXP", amt)
    put("Interest on loans", "FINANCE", pl["interest"])
    put("Round off", "OTHER_INCOME", -pl["round_off"])

    closing_stock = D(pl["closing_stock"])
    for name, amt in bs["assets"]:
        if name == "Closing stock":
            continue
        head = ("TRADE_RECEIVABLES" if name == "Sundry debtors" else "ST_LOANS_ADV" if re.search(r"tds|tcs|gst", name, re.I)
                else "OTHER_CUR_ASSETS" if name.startswith("Cheques") else "CASH_BANK")
        put(name, head, amt)
    for name, amt in bs["liabilities"]:
        head = ("LT_SECURED" if name.startswith("Loan:") else "TRADE_PAYABLES" if name == "Sundry creditors" else "OTHER_CUR_LIAB")
        put(name, head, -amt)
    drawings_fy = sum((c.amount for c in db.scalars(select(CapitalEntry).where(
        CapitalEntry.business_id == biz.id, CapitalEntry.date >= start, CapitalEntry.date <= end))
        if c.type.value == "DRAWINGS"), ZERO)
    capital_ledger = bs["capital"]["closing"] - pl["net_profit"] + drawings_fy
    put("Capital account", "CAPITAL", -capital_ledger)
    put("Drawings", "DRAWINGS", drawings_fy)
    if bs["difference"]:
        put("Difference as per books", "SUSPENSE", bs["difference"])
    return lines, closing_stock


def books_to_ledgers(db: Session, biz: Business, fy: str) -> tuple[list[dict], dict]:
    start = dt.date(int(fy[:4]), 4, 1)
    end = dt.date(start.year + 1, 3, 31)
    cy, cy_stock = _books_year(db, biz, start, end)
    py, py_stock = _books_year(db, biz, start.replace(year=start.year - 1), end.replace(year=end.year - 1))
    names = list(dict.fromkeys([*cy, *py]))
    ledgers = [dict(id=new_id(), name=n, group="MyBillSync books", head=(cy.get(n) or py.get(n))[0], partner=None,
                    cy=str(cy.get(n, ("", ZERO))[1]), py=str(py.get(n, ("", ZERO))[1])) for n in names]
    return ledgers, {"cy": str(cy_stock), "py": str(py_stock)}
