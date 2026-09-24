"""Bulk import from Excel (.xlsx) or CSV, with downloadable templates.

Each import runs inside one transaction: either every row is saved or none is.
A dry run validates everything and rolls back, so users can fix errors first.
Parties and items are matched (GSTIN / code / name) and updated if they exist,
which also makes the import a bulk-update tool.
"""

import csv
import datetime as dt
import io
import re
from collections import OrderedDict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError
from sqlalchemy import func, select

from ..deps import Ctx
from ..gst.constants import (
    GST_RATES,
    UQC,
    ExpenseKind,
    ItemType,
    PartyGstType,
    PartyType,
    PaymentMode,
    PaymentType,
    StockMoveType,
    VoucherType,
)
from ..gst.gstin import validate_gstin
from ..gst.states import STATES
from ..models import Account, ExpenseCategory, ExpenseItem, HsnCode, Item, Party, StockMovement, Voucher
from ..schemas import ItemIn, PartyIn, PaymentIn, VoucherIn, VoucherLineIn
from .payments import SETTLES, create_payment
from .vouchers import save_voucher

ZERO = Decimal("0")


@dataclass
class F:
    key: str
    header: str
    required: bool = False
    help: str = ""
    example: str = ""


# ---------------------------------------------------------------- column specs
PARTY_FIELDS = [
    F("name", "Party Name", True, "", "Sharma Traders"),
    F("type", "Party Type", False, "Customer / Supplier / Both", "Customer"),
    F("gstin", "GSTIN", False, "15 characters; fills state automatically", "27AAPFU0939F1ZV"),
    F("gst_type", "GST Type", False, "Registered / Composition / Unregistered / Consumer / SEZ / Overseas", "Registered"),
    F("phone", "Phone", False, "", "9876543210"),
    F("email", "Email", False, "", ""),
    F("state", "State", False, "State name or 2-digit code", "Maharashtra"),
    F("billing_address", "Billing Address", False, "", "12 MG Road"),
    F("city", "City", False, "", "Pune"),
    F("pincode", "Pincode", False, "", "411001"),
    F("shipping_address", "Shipping Address", False, "", ""),
    F("opening_balance", "Opening Balance", False, "Amount", "5000"),
    F("balance_type", "Balance Type", False, "Receivable (party owes you) / Payable (you owe party)", "Receivable"),
    F("credit_limit", "Credit Limit", False, "", ""),
]

ITEM_FIELDS = [
    F("name", "Item Name", True, "", "Steel Bottle 1L"),
    F("code", "Item Code", False, "SKU / barcode; used to match existing items", "SB-1L"),
    F("type", "Item Type", False, "Goods / Service", "Goods"),
    F("hsn_sac", "HSN/SAC", False, "4-8 digits", "7323"),
    F("unit", "Unit", False, "UQC code: " + ", ".join(UQC), "PCS"),
    F("category", "Category", False, "", "Bottles"),
    F("sale_price", "Sale Price", False, "", "450"),
    F("sale_price_tax_inclusive", "Sale Price Includes Tax", False, "Y / N", "N"),
    F("purchase_price", "Purchase Price", False, "", "300"),
    F("purchase_price_tax_inclusive", "Purchase Price Includes Tax", False, "Y / N", "N"),
    F("mrp", "MRP", False, "", "499"),
    F("gst_rate", "GST %", False, "One of " + ", ".join(str(r) for r in GST_RATES), "18"),
    F("cess_rate", "Cess %", False, "", "0"),
    F("opening_stock", "Opening Stock", False, "New items only", "50"),
    F("opening_stock_date", "Opening Stock Date", False, "DD/MM/YYYY", "01/04/2026"),
    F("low_stock_level", "Low Stock Alert", False, "", "10"),
    F("track_batch", "Track Batch", False, "Y / N", "N"),
    F("track_serial", "Track Serial", False, "Y / N", "N"),
    F("description", "Description", False, "", ""),
]

STOCK_FIELDS = [
    F("item", "Item Name", False, "Item name or code is required", "Steel Bottle 1L"),
    F("code", "Item Code", False, "", ""),
    F("date", "Date", True, "DD/MM/YYYY", "24/09/2026"),
    F("qty", "Quantity", True, "Positive adds stock, negative reduces stock", "25"),
    F("note", "Reason", False, "", "Physical count"),
]

DOC_HEAD = [
    F("doc_no", "Doc No", True, "Rows with the same Doc No become one document. Max 16 chars: letters, digits, / and -", "INV/26-27/0101"),
    F("date", "Date", True, "DD/MM/YYYY", "24/09/2026"),
]
PARTY_REF = [
    F("party", "Party Name", False, "Created if not found. Blank = cash / walk-in (must be fully paid)", "Karnataka Retail"),
    F("party_gstin", "Party GSTIN", False, "Used to match or create the party", ""),
    F("party_phone", "Party Phone", False, "", ""),
    F("pos", "Place of Supply", False, "State name or code; defaults to party state", ""),
]
LINE = [
    F("item", "Item Name", True, "Created if not found", "Steel Bottle 1L"),
    F("item_code", "Item Code", False, "", ""),
    F("hsn_sac", "HSN/SAC", False, "", "7323"),
    F("qty", "Qty", True, "", "3"),
    F("unit", "Unit", False, "", "PCS"),
    F("rate", "Rate", True, "Price per unit", "450"),
    F("tax_inclusive", "Rate Includes Tax", False, "Y / N", "N"),
    F("discount_pct", "Discount %", False, "", "0"),
    F("gst_rate", "GST %", False, "", "18"),
    F("cess_rate", "Cess %", False, "", "0"),
    F("batch_no", "Batch No", False, "", ""),
    F("expiry_date", "Expiry Date", False, "DD/MM/YYYY", ""),
    F("serial_nos", "Serial Nos", False, "Comma separated", ""),
]
DOC_TAIL = [
    F("due_date", "Due Date", False, "DD/MM/YYYY", ""),
    F("notes", "Notes", False, "", ""),
]
PAY = [
    F("amount_paid", "Amount Paid", False, "Paid/received now (first row of the document). Write FULL for the whole bill", ""),
    F("payment_mode", "Payment Mode", False, "Cash / Bank / UPI / Cheque / Card", "Cash"),
    F("account", "Account", False, "Cash / bank account name; default Cash in Hand", ""),
]
TCS = [F("tcs_rate", "TCS %", False, "Tax collected at source, if applicable", "")]
PURCHASE_EXTRA = [
    F("supplier_bill_no", "Supplier Bill No", False, "", "W-77"),
    F("supplier_bill_date", "Supplier Bill Date", False, "DD/MM/YYYY", ""),
    F("gst_charged", "GST Charged", False, "Y / N — did the supplier charge GST? Default: yes if supplier has GSTIN", ""),
    F("reverse_charge", "Reverse Charge", False, "Y / N", "N"),
]
RETURN_EXTRA = [
    F("original_doc", "Original Doc No", False, "Invoice/bill this note is against", ""),
    F("reason", "Reason", False, "", "Sales return"),
]
EXPENSE_FIELDS = DOC_HEAD + [
    F("category", "Category", True, "Created if not found", "Rent"),
    F("category_type", "Category Type", False, "Direct / Indirect (for new categories)", "Indirect"),
    F("item", "Expense Item", True, "Created if not found", "Shop rent"),
    F("hsn_sac", "HSN/SAC", False, "", ""),
    F("qty", "Qty", False, "Default 1", "1"),
    F("rate", "Amount", True, "", "25000"),
    F("gst_rate", "GST %", False, "", "0"),
    F("party", "Paid To (Party)", False, "Optional supplier; blank = paid directly", ""),
    F("party_gstin", "Party GSTIN", False, "", ""),
    F("gst_charged", "GST Bill", False, "Y / N — claim input credit", "N"),
] + PAY + [F("notes", "Notes", False, "", "")]

PAYMENT_FIELDS = [
    F("date", "Date", True, "DD/MM/YYYY", "24/09/2026"),
    F("party", "Party Name", True, "", "Karnataka Retail"),
    F("party_gstin", "Party GSTIN", False, "", ""),
    F("amount", "Amount", True, "", "5000"),
    F("tds", "TDS Amount", False, "", "0"),
    F("mode", "Mode", False, "Cash / Bank / UPI / Cheque / Card", "UPI"),
    F("account", "Account", False, "Cash / bank account name", ""),
    F("reference", "Reference", False, "UTR / cheque no.", ""),
    F("against", "Against Doc No", False, "Bill to settle first; others settle oldest first", ""),
    F("notes", "Notes", False, "", ""),
]

HSN_FIELDS = [
    F("code", "HSN/SAC Code", True, "4-8 digits", "7323"),
    F("description", "Description", False, "", "Table, kitchen or other household articles of iron or steel"),
    F("gst_rate", "GST %", True, "", "18"),
    F("cess_rate", "Cess %", False, "", "0"),
    F("effective_from", "Effective From", False, "DD/MM/YYYY", "22/09/2025"),
]

EXPENSE_ITEM_FIELDS = [
    F("name", "Expense Item", True, "", "Petrol"),
    F("category", "Category", False, "Created if not found", "Petrol & Fuel"),
    F("hsn_sac", "HSN/SAC", False, "", ""),
    F("rate", "Default Amount", False, "", ""),
    F("gst_rate", "GST %", False, "", "0"),
]

VOUCHER_ENTITIES = {
    "sales": (VoucherType.SALE, DOC_HEAD + PARTY_REF + LINE + DOC_TAIL + PAY + TCS, "Sale invoices"),
    "estimates": (VoucherType.ESTIMATE, DOC_HEAD + PARTY_REF + LINE + DOC_TAIL, "Estimates / quotations"),
    "sale-orders": (VoucherType.SALE_ORDER, DOC_HEAD + PARTY_REF + LINE + DOC_TAIL, "Sale orders"),
    "delivery-challans": (VoucherType.DELIVERY_CHALLAN, DOC_HEAD + PARTY_REF + LINE + DOC_TAIL, "Delivery challans"),
    "credit-notes": (VoucherType.SALE_RETURN, DOC_HEAD + PARTY_REF + RETURN_EXTRA + LINE + DOC_TAIL + PAY, "Credit notes"),
    "purchases": (VoucherType.PURCHASE, DOC_HEAD + PARTY_REF + PURCHASE_EXTRA + LINE + DOC_TAIL + PAY + TCS, "Purchase bills"),
    "purchase-orders": (VoucherType.PURCHASE_ORDER, DOC_HEAD + PARTY_REF + LINE + DOC_TAIL, "Purchase orders"),
    "debit-notes": (VoucherType.PURCHASE_RETURN, DOC_HEAD + PARTY_REF + RETURN_EXTRA + LINE + DOC_TAIL + PAY, "Debit notes"),
    "expenses": (VoucherType.EXPENSE, EXPENSE_FIELDS, "Expenses"),
}

ENTITIES: dict[str, tuple[list[F], str]] = {
    "parties": (PARTY_FIELDS, "Parties"),
    "items": (ITEM_FIELDS, "Items"),
    "stock": (STOCK_FIELDS, "Stock adjustments"),
    "payments-in": (PAYMENT_FIELDS, "Payments received"),
    "payments-out": (PAYMENT_FIELDS, "Payments made"),
    "hsn": (HSN_FIELDS, "HSN/SAC master"),
    "expense-items": (EXPENSE_ITEM_FIELDS, "Expense items"),
    **{k: (v[1], v[2]) for k, v in VOUCHER_ENTITIES.items()},
}


def spec(entity: str) -> tuple[list[F], str]:
    if entity not in ENTITIES:
        raise HTTPException(404, "Unknown import type")
    return ENTITIES[entity]


# ---------------------------------------------------------------- template
def template(entity: str) -> bytes:
    fields, title = spec(entity)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    head_fill = PatternFill("solid", fgColor="1F65BB")
    for i, f in enumerate(fields, 1):
        c = ws.cell(row=1, column=i, value=f.header + (" *" if f.required else ""))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row=2, column=i, value=f.example or None)
        ws.column_dimensions[c.column_letter].width = max(14, len(f.header) + 4)
    ws.freeze_panes = "A2"
    info = wb.create_sheet("Instructions")
    info.append([f"{title} — import template"])
    info["A1"].font = Font(bold=True, size=13)
    info.append(["Fill the 'Data' sheet from row 2 (delete the example row). Columns marked * are required."])
    info.append(["Dates: DD/MM/YYYY. Amounts: plain numbers. Y/N columns accept Y, Yes, N, No."])
    info.append([])
    info.append(["Column", "Required", "Notes"])
    for c in info[5]:
        c.font = Font(bold=True)
    for f in fields:
        info.append([f.header, "Yes" if f.required else "", f.help])
    info.column_dimensions["A"].width = 26
    info.column_dimensions["C"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- parsing helpers
def _norm(h) -> str:
    return re.sub(r"[^a-z0-9%]", "", str(h or "").lower().replace("*", ""))


def read_rows(filename: str, content: bytes, fields: list[F]) -> list[tuple[int, dict]]:
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        raw = list(csv.reader(io.StringIO(text)))
    else:
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise HTTPException(400, "Could not read the file — upload the .xlsx template or a .csv") from e
        ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    if not raw:
        raise HTTPException(400, "The file is empty")
    by_header = {_norm(f.header): f.key for f in fields}
    cols = [by_header.get(_norm(h)) for h in raw[0]]
    if not any(cols):
        raise HTTPException(400, "Column headings do not match the template — download the template and use it")
    rows = []
    for idx, r in enumerate(raw[1:], start=2):
        rec = {k: v for k, v in zip(cols, r) if k and v not in (None, "")}
        rec = {k: (v.strip() if isinstance(v, str) else v) for k, v in rec.items()}
        rec = {k: v for k, v in rec.items() if v != ""}
        if rec:
            rows.append((idx, rec))
    if not rows:
        raise HTTPException(400, "No data rows found")
    return rows


class RowError(Exception):
    pass


def num(v, field: str, default=None) -> Decimal | None:
    if v is None:
        return default
    try:
        return Decimal(str(v).replace(",", "").replace("₹", "").strip())
    except InvalidOperation as e:
        raise RowError(f"{field}: '{v}' is not a number") from e


def date(v, field: str) -> dt.date | None:
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise RowError(f"{field}: '{s}' is not a valid date (use DD/MM/YYYY)")


def yes(v) -> bool | None:
    if v is None:
        return None
    return str(v).strip().lower() in ("y", "yes", "true", "1")


def state_code(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if s.isdigit():
        s = s.zfill(2)
        if s in STATES:
            return s
    for code, name in STATES.items():
        if name.lower() == s.lower():
            return code
    raise RowError(f"Unknown state '{v}'")


def enum_of(enum_cls, v, default):
    if v is None:
        return default
    key = str(v).strip().upper().replace(" ", "_")
    aliases = {"PAYABLE": "PAYABLE", "RECEIVABLE": "RECEIVABLE"}
    key = aliases.get(key, key)
    try:
        return enum_cls(key)
    except ValueError as e:
        raise RowError(f"'{v}' is not valid; use one of: {', '.join(x.value.title() for x in enum_cls)}") from e


def _pyd_err(e: ValidationError) -> str:
    return "; ".join(
        f"{'.'.join(str(x) for x in err['loc'] if x not in ('body',))}: {err['msg'].replace('Value error, ', '')}"
        for err in e.errors())


# ---------------------------------------------------------------- lookups (with in-run creation)
class Resolver:
    def __init__(self, ctx: Ctx):
        self.ctx = ctx
        self.db = ctx.db
        self.created = {"parties": 0, "items": 0, "categories": 0, "expense_items": 0}

    def party(self, name, gstin, phone=None, ptype=PartyType.CUSTOMER, state=None, create=True) -> Party | None:
        db, bid = self.db, self.ctx.bid
        p = None
        if gstin:
            gstin = str(gstin).upper()
            p = db.scalar(select(Party).where(Party.business_id == bid, Party.gstin == gstin))
        if p is None and name:
            p = db.scalar(select(Party).where(Party.business_id == bid, func.lower(Party.name) == str(name).lower()))
        if p is None and (name or gstin) and create:
            if gstin and validate_gstin(gstin):
                raise RowError(f"Party GSTIN: {validate_gstin(gstin)}")
            p = Party(business_id=bid, name=str(name or gstin), type=ptype, gstin=gstin or None,
                      gst_type=PartyGstType.REGISTERED if gstin else (
                          PartyGstType.CONSUMER if ptype == PartyType.CUSTOMER else PartyGstType.UNREGISTERED),
                      state_code=gstin[:2] if gstin else state, phone=str(phone) if phone else None,
                      pan=gstin[2:12] if gstin else None)
            db.add(p)
            db.flush()
            self.created["parties"] += 1
        return p

    def item(self, name, code, hsn, unit, rate, gst_rate, sale: bool) -> Item:
        db, bid = self.db, self.ctx.bid
        it = None
        if code:
            it = db.scalar(select(Item).where(Item.business_id == bid, Item.code == str(code)))
        if it is None and name:
            it = db.scalar(select(Item).where(Item.business_id == bid, func.lower(Item.name) == str(name).lower()))
        if it is None:
            hsn = str(hsn) if hsn else None
            it = Item(business_id=bid, name=str(name), code=str(code) if code else None, hsn_sac=hsn,
                      type=ItemType.SERVICE if hsn and hsn.startswith("99") else ItemType.GOODS,
                      unit=(str(unit).upper() if unit and str(unit).upper() in UQC else "NOS"),
                      sale_price=rate if sale else ZERO, purchase_price=ZERO if sale else rate,
                      gst_rate=gst_rate or ZERO)
            db.add(it)
            db.flush()
            self.created["items"] += 1
        return it

    def category(self, name, kind=None) -> ExpenseCategory:
        c = self.db.scalar(select(ExpenseCategory).where(ExpenseCategory.business_id == self.ctx.bid,
                                                         func.lower(ExpenseCategory.name) == str(name).lower()))
        if c is None:
            c = ExpenseCategory(business_id=self.ctx.bid, name=str(name),
                                kind=enum_of(ExpenseKind, kind, ExpenseKind.INDIRECT))
            self.db.add(c)
            self.db.flush()
            self.created["categories"] += 1
        return c

    def expense_item(self, name, category_id, hsn, rate, gst_rate) -> ExpenseItem:
        e = self.db.scalar(select(ExpenseItem).where(ExpenseItem.business_id == self.ctx.bid,
                                                     func.lower(ExpenseItem.name) == str(name).lower()))
        if e is None:
            e = ExpenseItem(business_id=self.ctx.bid, name=str(name), category_id=category_id,
                            hsn_sac=str(hsn) if hsn else None, rate=rate or ZERO, gst_rate=gst_rate or ZERO)
            self.db.add(e)
            self.db.flush()
            self.created["expense_items"] += 1
        return e

    def account(self, name) -> str | None:
        if not name:
            return None
        a = self.db.scalar(select(Account).where(Account.business_id == self.ctx.bid,
                                                 func.lower(Account.name) == str(name).lower()))
        if not a:
            raise RowError(f"Account '{name}' not found — create it under Cash & Bank first")
        return a.id

    def voucher_by_number(self, number, types) -> Voucher | None:
        if not number:
            return None
        v = self.db.scalar(select(Voucher).where(Voucher.business_id == self.ctx.bid, Voucher.type.in_(types),
                                                 Voucher.number == str(number)))
        if not v:
            raise RowError(f"Document '{number}' not found")
        return v


# ---------------------------------------------------------------- importers
def _import_parties(ctx: Ctx, rows, res: Resolver):
    created = updated = 0
    for n, r in rows:
        try:
            gstin = str(r["gstin"]).upper() if r.get("gstin") else None
            existing = res.party(r.get("name"), gstin, create=False)
            bal = num(r.get("opening_balance"), "Opening Balance", ZERO)
            if str(r.get("balance_type", "")).lower().startswith("pay"):
                bal = -abs(bal)
            data = PartyIn(
                name=r.get("name"), type=enum_of(PartyType, r.get("type"), PartyType.CUSTOMER),
                gst_type=enum_of(PartyGstType, r.get("gst_type"),
                                 PartyGstType.REGISTERED if gstin else PartyGstType.UNREGISTERED),
                gstin=gstin, phone=str(r["phone"]) if r.get("phone") else None, email=r.get("email"),
                state_code=state_code(r.get("state")), billing_address=r.get("billing_address"),
                city=r.get("city"), pincode=str(r["pincode"]) if r.get("pincode") else None,
                shipping_address=r.get("shipping_address"), opening_balance=bal,
                credit_limit=num(r.get("credit_limit"), "Credit Limit"))
            if existing:
                for k, v in data.model_dump().items():
                    setattr(existing, k, v)
                updated += 1
            else:
                ctx.db.add(Party(business_id=ctx.bid, **data.model_dump()))
                created += 1
            ctx.db.flush()
        except ValidationError as e:
            yield n, _pyd_err(e)
        except RowError as e:
            yield n, str(e)
    res.created["parties"] += created
    res.summary = f"{created} parties created, {updated} updated"


def _import_items(ctx: Ctx, rows, res: Resolver):
    created = updated = 0
    for n, r in rows:
        try:
            existing = None
            if r.get("code"):
                existing = ctx.db.scalar(select(Item).where(Item.business_id == ctx.bid, Item.code == str(r["code"])))
            if existing is None and r.get("name"):
                existing = ctx.db.scalar(select(Item).where(Item.business_id == ctx.bid,
                                                            func.lower(Item.name) == str(r["name"]).lower()))
            base = {} if existing is None else {
                k: getattr(existing, k) for k in ItemIn.model_fields if hasattr(existing, k)}
            vals = dict(
                name=r.get("name"), code=str(r["code"]) if r.get("code") else base.get("code"),
                type=enum_of(ItemType, r.get("type"), base.get("type", ItemType.GOODS)),
                hsn_sac=str(r["hsn_sac"]) if r.get("hsn_sac") else base.get("hsn_sac"),
                unit=str(r.get("unit") or base.get("unit") or "NOS").upper(),
                category=r.get("category", base.get("category")),
                sale_price=num(r.get("sale_price"), "Sale Price", base.get("sale_price", ZERO)),
                purchase_price=num(r.get("purchase_price"), "Purchase Price", base.get("purchase_price", ZERO)),
                mrp=num(r.get("mrp"), "MRP", base.get("mrp")),
                gst_rate=num(r.get("gst_rate"), "GST %", base.get("gst_rate", ZERO)),
                cess_rate=num(r.get("cess_rate"), "Cess %", base.get("cess_rate", ZERO)),
                low_stock_level=num(r.get("low_stock_level"), "Low Stock Alert", base.get("low_stock_level")),
                description=r.get("description", base.get("description")),
            )
            for k in ("sale_price_tax_inclusive", "purchase_price_tax_inclusive", "track_batch", "track_serial"):
                y = yes(r.get(k))
                vals[k] = y if y is not None else base.get(k, False)
            data = ItemIn(**vals)
            fields = data.model_dump(exclude={"opening_stock", "opening_stock_date"})
            if existing:
                for k, v in fields.items():
                    setattr(existing, k, v)
                updated += 1
            else:
                it = Item(business_id=ctx.bid, **fields)
                ctx.db.add(it)
                ctx.db.flush()
                qty = num(r.get("opening_stock"), "Opening Stock", ZERO)
                if qty and it.type == ItemType.GOODS:
                    ctx.db.add(StockMovement(business_id=ctx.bid, item_id=it.id,
                                             date=date(r.get("opening_stock_date"), "Opening Stock Date") or dt.date.today(),
                                             type=StockMoveType.OPENING, qty=qty, rate=it.purchase_price,
                                             note="Opening stock (import)"))
                created += 1
            ctx.db.flush()
        except ValidationError as e:
            yield n, _pyd_err(e)
        except RowError as e:
            yield n, str(e)
    res.summary = f"{created} items created, {updated} updated"


def _import_stock(ctx: Ctx, rows, res: Resolver):
    count = 0
    for n, r in rows:
        try:
            it = None
            if r.get("code"):
                it = ctx.db.scalar(select(Item).where(Item.business_id == ctx.bid, Item.code == str(r["code"])))
            if it is None and r.get("item"):
                it = ctx.db.scalar(select(Item).where(Item.business_id == ctx.bid,
                                                      func.lower(Item.name) == str(r["item"]).lower()))
            if it is None:
                raise RowError("Item not found")
            if it.type != ItemType.GOODS:
                raise RowError("Stock is not tracked for services")
            qty = num(r.get("qty"), "Quantity")
            if not qty:
                raise RowError("Quantity is required")
            ctx.db.add(StockMovement(business_id=ctx.bid, item_id=it.id, date=date(r.get("date"), "Date"),
                                     type=StockMoveType.ADJUSTMENT, qty=qty, rate=it.purchase_price,
                                     note=r.get("note") or "Imported adjustment"))
            count += 1
        except RowError as e:
            yield n, str(e)
    res.summary = f"{count} stock adjustments"


def _import_payments(ctx: Ctx, rows, res: Resolver, ptype: PaymentType):
    count = 0
    for n, r in rows:
        try:
            party = res.party(r.get("party"), r.get("party_gstin"), create=False)
            if not party:
                raise RowError("Party not found")
            against = res.voucher_by_number(r.get("against"), SETTLES[ptype])
            data = PaymentIn(type=ptype, date=date(r.get("date"), "Date"), party_id=party.id,
                             amount=num(r.get("amount"), "Amount", ZERO), tds_amount=num(r.get("tds"), "TDS", ZERO),
                             mode=enum_of(PaymentMode, r.get("mode"), PaymentMode.CASH),
                             account_id=res.account(r.get("account")), reference=r.get("reference"),
                             notes=r.get("notes"), voucher_id=against.id if against else None)
            create_payment(ctx, data)
            count += 1
        except ValidationError as e:
            yield n, _pyd_err(e)
        except (RowError, HTTPException) as e:
            yield n, getattr(e, "detail", None) or str(e)
    res.summary = f"{count} payments"


def _import_hsn(ctx: Ctx, rows, res: Resolver):
    created = updated = 0
    for n, r in rows:
        try:
            code = re.sub(r"\D", "", str(r.get("code", "")))
            if not 4 <= len(code) <= 8:
                raise RowError("HSN/SAC must be 4-8 digits")
            rate = num(r.get("gst_rate"), "GST %")
            if rate not in GST_RATES:
                raise RowError(f"GST % must be one of {', '.join(str(x) for x in GST_RATES)}")
            h = ctx.db.scalar(select(HsnCode).where(HsnCode.business_id == ctx.bid, HsnCode.code == code))
            if h is None:
                h = HsnCode(business_id=ctx.bid, code=code)
                ctx.db.add(h)
                created += 1
            else:
                updated += 1
            h.description = r.get("description") or h.description
            h.gst_rate = rate
            h.cess_rate = num(r.get("cess_rate"), "Cess %", ZERO)
            h.effective_from = date(r.get("effective_from"), "Effective From")
            ctx.db.flush()
        except RowError as e:
            yield n, str(e)
    res.summary = f"{created} codes added, {updated} updated"


def _import_expense_items(ctx: Ctx, rows, res: Resolver):
    count = 0
    for n, r in rows:
        try:
            cat = res.category(r["category"]) if r.get("category") else None
            rate = num(r.get("gst_rate"), "GST %", ZERO)
            if rate not in GST_RATES:
                raise RowError("Invalid GST %")
            e = res.expense_item(r.get("name"), cat.id if cat else None, r.get("hsn_sac"),
                                 num(r.get("rate"), "Default Amount", ZERO), rate)
            e.category_id = cat.id if cat else e.category_id
            count += 1
        except RowError as e:
            yield n, str(e)
    res.summary = f"{count} expense items"


def _import_vouchers(ctx: Ctx, rows, res: Resolver, vtype: VoucherType):
    groups: "OrderedDict[str, list[tuple[int, dict]]]" = OrderedDict()
    for n, r in rows:
        key = str(r.get("doc_no") or "").strip()
        if not key:
            yield n, "Doc No is required"
            continue
        groups.setdefault(key, []).append((n, r))

    outward = vtype in (VoucherType.SALE, VoucherType.SALE_RETURN, VoucherType.ESTIMATE,
                        VoucherType.SALE_ORDER, VoucherType.DELIVERY_CHALLAN)
    ptype = PartyType.CUSTOMER if outward else PartyType.SUPPLIER
    count = 0
    for doc_no, grp in groups.items():
        first_row, h = grp[0]
        try:
            pos = state_code(h.get("pos"))
            party = res.party(h.get("party"), h.get("party_gstin"), h.get("party_phone"), ptype, pos)
            lines = []
            category = None
            if vtype == VoucherType.EXPENSE:
                category = res.category(h.get("category") or "Miscellaneous", h.get("category_type"))
            for _, r in grp:
                rate = num(r.get("rate"), "Rate")
                if rate is None:
                    raise RowError("Rate / Amount is required")
                gst_rate = num(r.get("gst_rate"), "GST %", None)
                if vtype == VoucherType.EXPENSE:
                    e = res.expense_item(r.get("item"), category.id, r.get("hsn_sac"), rate, gst_rate)
                    lines.append(VoucherLineIn(expense_item_id=e.id, name=e.name, hsn_sac=r.get("hsn_sac") or e.hsn_sac,
                                               qty=num(r.get("qty"), "Qty", Decimal("1")), rate=rate,
                                               gst_rate=gst_rate if gst_rate is not None else e.gst_rate))
                else:
                    it = res.item(r.get("item"), r.get("item_code"), r.get("hsn_sac"), r.get("unit"), rate,
                                  gst_rate, sale=outward)
                    lines.append(VoucherLineIn(
                        item_id=it.id, name=it.name, hsn_sac=str(r["hsn_sac"]) if r.get("hsn_sac") else it.hsn_sac,
                        unit=str(r.get("unit") or it.unit).upper(), qty=num(r.get("qty"), "Qty"), rate=rate,
                        tax_inclusive=bool(yes(r.get("tax_inclusive"))),
                        discount_pct=num(r.get("discount_pct"), "Discount %", ZERO),
                        gst_rate=gst_rate if gst_rate is not None else it.gst_rate,
                        cess_rate=num(r.get("cess_rate"), "Cess %", it.cess_rate),
                        batch_no=str(r["batch_no"]) if r.get("batch_no") else None,
                        expiry_date=date(r.get("expiry_date"), "Expiry Date"),
                        serial_nos=str(r["serial_nos"]) if r.get("serial_nos") else None))
            paid_raw = h.get("amount_paid")
            full = (str(paid_raw).strip().upper() == "FULL") or (party is None and paid_raw is None
                                                                  and vtype in (VoucherType.SALE, VoucherType.EXPENSE,
                                                                                VoucherType.PURCHASE))
            original = None
            if h.get("original_doc"):
                original = res.voucher_by_number(h["original_doc"], [VoucherType.SALE if outward else VoucherType.PURCHASE])
            data = VoucherIn(
                type=vtype, number=doc_no, date=date(h.get("date"), "Date"), due_date=date(h.get("due_date"), "Due Date"),
                party_id=party.id if party else None, party_phone=str(h["party_phone"]) if h.get("party_phone") else None,
                place_of_supply=pos, tax_applicable=yes(h.get("gst_charged")),
                reverse_charge=bool(yes(h.get("reverse_charge"))),
                supplier_invoice_no=str(h["supplier_bill_no"]) if h.get("supplier_bill_no") else None,
                supplier_invoice_date=date(h.get("supplier_bill_date"), "Supplier Bill Date"),
                original_voucher_id=original.id if original else None, reason=h.get("reason"),
                notes=h.get("notes"), expense_category_id=category.id if category else None,
                tcs_rate=num(h.get("tcs_rate"), "TCS %", ZERO), lines=lines,
                amount_paid=ZERO if full else num(paid_raw, "Amount Paid", ZERO), fully_paid=full,
                payment_mode=enum_of(PaymentMode, h.get("payment_mode"), PaymentMode.CASH),
                payment_account_id=res.account(h.get("account")),
            )
            with ctx.db.begin_nested():
                save_voucher(ctx, data)
            count += 1
        except ValidationError as e:
            yield first_row, f"{doc_no}: {_pyd_err(e)}"
        except RowError as e:
            yield first_row, f"{doc_no}: {e}"
        except HTTPException as e:
            yield first_row, f"{doc_no}: {e.detail}"
    res.summary = f"{count} documents"


def run_import(ctx: Ctx, entity: str, filename: str, content: bytes, dry_run: bool) -> dict:
    fields, title = spec(entity)
    rows = read_rows(filename, content, fields)
    res = Resolver(ctx)
    res.summary = ""
    missing_required = []
    for n, r in rows:
        miss = [f.header for f in fields if f.required and f.key not in r]
        if miss and entity not in VOUCHER_ENTITIES:
            missing_required.append({"row": n, "message": f"Missing: {', '.join(miss)}"})
    if entity in VOUCHER_ENTITIES:
        gen = _import_vouchers(ctx, rows, res, VOUCHER_ENTITIES[entity][0])
    else:
        gen = {
            "parties": lambda: _import_parties(ctx, rows, res),
            "items": lambda: _import_items(ctx, rows, res),
            "stock": lambda: _import_stock(ctx, rows, res),
            "payments-in": lambda: _import_payments(ctx, rows, res, PaymentType.IN),
            "payments-out": lambda: _import_payments(ctx, rows, res, PaymentType.OUT),
            "hsn": lambda: _import_hsn(ctx, rows, res),
            "expense-items": lambda: _import_expense_items(ctx, rows, res),
        }[entity]()
    errors = missing_required + [{"row": n, "message": m} for n, m in gen]
    ok = not errors
    if ok and not dry_run:
        ctx.db.commit()
    else:
        ctx.db.rollback()
    return {"title": title, "rows": len(rows), "ok": ok, "dry_run": dry_run, "saved": ok and not dry_run,
            "summary": res.summary, "created": res.created, "errors": sorted(errors, key=lambda e: e["row"])[:500]}
