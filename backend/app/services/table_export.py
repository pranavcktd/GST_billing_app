"""Excel (.xlsx) and PDF files from any table shown in the app.

Input is the report shape used across the app:
    {title, subtitle, summary: [{label, value, type}],
     sections: [{title, columns: [{key, label, type}], rows: [...], total: {...}, note}]}
Column types: text, money, qty, date, pct, int. Rows may carry `_style` (bold | head | sub).
"""

import datetime as dt
import io
from decimal import Decimal, InvalidOperation
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

NUMERIC = {"money", "qty", "pct", "int"}
MAX_ROWS = 50000


def _num(v):
    if v is None or v == "":
        return None
    try:
        return Decimal(str(v))
    except (InvalidOperation, ValueError):
        return None


def indian(n: Decimal, places: int = 2) -> str:
    """12,34,567.89 grouping."""
    neg = n < 0
    s = f"{abs(n):.{places}f}"
    whole, _, frac = s.partition(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        whole = ",".join(groups + [tail])
    return ("-" if neg else "") + whole + (f".{frac}" if frac else "")


def fmt(v, t: str) -> str:
    if v is None or v == "":
        return ""
    if t in NUMERIC:
        n = _num(v)
        if n is None:
            return str(v)
        if t == "money":
            return indian(n)
        if t == "pct":
            return f"{n.normalize():f}%"
        if t == "int":
            return f"{int(n)}"
        return f"{n.normalize():f}" if n == n.to_integral() else f"{n:.3f}".rstrip("0").rstrip(".")
    if t == "date":
        try:
            return dt.date.fromisoformat(str(v)[:10]).strftime("%d-%m-%Y")
        except ValueError:
            return str(v)
    return str(v)


def _meta_lines(doc: dict, business: str | None) -> list[str]:
    lines = [x for x in (business, doc.get("subtitle")) if x]
    lines.append(f"Generated on {dt.datetime.now():%d-%m-%Y %H:%M}")
    return lines


# ================================================================ Excel
def to_xlsx(doc: dict, business: str | None = None, brand: str = "SmartHisab") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (doc.get("title") or "Report")[:31].replace("/", "-").replace(":", "-") or "Report"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="1F65BB")
    thin = Side(style="thin", color="D1D5DB")
    ws.append([doc.get("title") or "Report"])
    ws["A1"].font = Font(bold=True, size=14)
    for line in _meta_lines(doc, business):
        ws.append([line])
    for s in doc.get("summary") or []:
        ws.append([s.get("label"), float(_num(s.get("value"))) if s.get("type") in NUMERIC and _num(s.get("value")) is not None else s.get("value")])
    widths: dict[int, int] = {}
    for sec in doc.get("sections") or []:
        ws.append([])
        if sec.get("title"):
            ws.append([sec["title"]])
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        cols = sec.get("columns") or []
        ws.append([c.get("label", "") for c in cols])
        r = ws.max_row
        for i, _ in enumerate(cols, 1):
            cell = ws.cell(r, i)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in (sec.get("rows") or [])[:MAX_ROWS]:
            values = []
            for c in cols:
                v = row.get(c["key"])
                t = c.get("type", "text")
                if t in NUMERIC and _num(v) is not None:
                    values.append(float(_num(v)))
                elif t == "date" and v:
                    try:
                        values.append(dt.date.fromisoformat(str(v)[:10]))
                    except ValueError:
                        values.append(str(v))
                else:
                    values.append(v if v is None or isinstance(v, (int, float)) else str(v))
            ws.append(values)
            rr = ws.max_row
            for i, c in enumerate(cols, 1):
                cell = ws.cell(rr, i)
                t = c.get("type", "text")
                cell.border = Border(bottom=thin)
                if t == "money":
                    cell.number_format = "#,##,##0.00;[Red]-#,##,##0.00"
                elif t == "date":
                    cell.number_format = "dd-mm-yyyy"
                elif t == "pct":
                    cell.number_format = '0.00"%"'
                if row.get("_style") in ("bold", "head"):
                    cell.font = bold
                widths[i] = max(widths.get(i, 8), min(60, len(fmt(row.get(c["key"]), t)) + 2))
        if sec.get("total"):
            ws.append([("Total" if i == 0 else (float(_num(sec["total"].get(c["key"]))) if _num(sec["total"].get(c["key"])) is not None else None))
                       for i, c in enumerate(cols)])
            for i, c in enumerate(cols, 1):
                cell = ws.cell(ws.max_row, i)
                cell.font = bold
                if c.get("type") == "money":
                    cell.number_format = "#,##,##0.00;[Red]-#,##,##0.00"
        for i, c in enumerate(cols, 1):
            widths[i] = max(widths.get(i, 8), min(40, len(c.get("label", "")) + 2))
        if sec.get("note"):
            ws.append([sec["note"]])
            ws.cell(ws.max_row, 1).font = Font(italic=True, color="6B7280")
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.append([])
    ws.append([f"Generated with {brand}. Please review before use in any filing."])
    ws.cell(ws.max_row, 1).font = Font(italic=True, color="9CA3AF", size=8)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ================================================================ PDF
def to_pdf(doc: dict, business: str | None = None, brand: str = "SmartHisab") -> bytes:
    sections = doc.get("sections") or []
    max_cols = max((len(s.get("columns") or []) for s in sections), default=1)
    size = landscape(A4) if max_cols > 6 else A4
    buf = io.BytesIO()
    title = doc.get("title") or "Report"

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#9CA3AF"))
        canvas.drawString(12 * mm, 8 * mm, f"{title} · Generated with {brand} — please review before use in any filing")
        canvas.drawRightString(size[0] - 12 * mm, 8 * mm, f"Page {d.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(buf, pagesize=size, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
                            bottomMargin=14 * mm, title=title, author=brand)
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=14, leading=17, spaceAfter=2)
    meta = ParagraphStyle("meta", fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#6B7280"))
    h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=8, spaceAfter=4)
    font_size = 8 if max_cols <= 8 else 7 if max_cols <= 12 else 6
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=font_size, leading=font_size + 2)
    cell_r = ParagraphStyle("cellr", parent=cell, alignment=2)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    head_r = ParagraphStyle("headr", parent=head, alignment=2)
    story = [Paragraph(escape(title), h1)] + [Paragraph(escape(x), meta) for x in _meta_lines(doc, business)]
    if doc.get("summary"):
        story.append(Spacer(1, 4))
        story.append(Paragraph(" &nbsp;·&nbsp; ".join(f"<b>{escape(str(s.get('label')))}:</b> {escape(fmt(s.get('value'), s.get('type', 'text')))}"
                                                        for s in doc["summary"]), meta))
    width = size[0] - 24 * mm
    for sec in sections:
        cols = sec.get("columns") or []
        if not cols:
            continue
        if sec.get("title"):
            story.append(Paragraph(escape(sec["title"]), h2))
        else:
            story.append(Spacer(1, 8))
        rows = (sec.get("rows") or [])[:MAX_ROWS]
        if not rows:
            story.append(Paragraph("No records.", meta))
            continue
        # column widths: text columns get more room
        weights = [3.0 if c.get("type", "text") == "text" else 1.4 if c.get("type") == "money" else 1.1 for c in cols]
        widths = [width * w / sum(weights) for w in weights]
        data = [[Paragraph(escape(c.get("label", "")), head_r if c.get("type") in NUMERIC else head) for c in cols]]
        styles = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F65BB")),
                  ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
                  ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
        for i, row in enumerate(rows, 1):
            data.append([Paragraph(escape(fmt(row.get(c["key"]), c.get("type", "text"))), cell_r if c.get("type") in NUMERIC else cell)
                         for c in cols])
            if row.get("_style") in ("bold", "head"):
                styles.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F3F4F6")))
        if sec.get("total"):
            data.append([Paragraph("<b>Total</b>" if j == 0 else f"<b>{escape(fmt(sec['total'].get(c['key']), c.get('type', 'text')))}</b>",
                                   cell_r if c.get("type") in NUMERIC and j else cell) for j, c in enumerate(cols)])
            styles.append(("LINEABOVE", (0, len(data) - 1), (-1, len(data) - 1), 0.8, colors.black))
        t = LongTable(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle(styles))
        story.append(t)
        if sec.get("note"):
            story.append(Paragraph(escape(sec["note"]), meta))
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
