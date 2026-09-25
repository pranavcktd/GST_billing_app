"""Cash & bank, expenses, orders/challans, TDS/TCS, cheques, loans, reports, imports, backup/restore."""

import csv
import io

from openpyxl import load_workbook

from app.gst.gstin import gstin_check_char
from tests.test_api_flow import make_business, signup


def gstin(state: str, pan: str) -> str:
    return f"{state}{pan}1Z" + gstin_check_char(f"{state}{pan}1Z")


def post(client, h, url, body, code=201):
    r = client.post(url, headers=h, json=body)
    assert r.status_code == code, r.text
    return r.json()


def build_books(client):
    """A realistic month of activity touching every module."""
    h = make_business(client, signup(client))
    accounts = client.get("/api/accounts", headers=h).json()
    cash = next(a for a in accounts if a["is_default_cash"])
    bank = post(client, h, "/api/accounts", {"type": "BANK", "name": "HDFC Current", "opening_balance": 10000})

    post(client, h, "/api/capital", {"date": "2026-04-01", "type": "INTRODUCED", "amount": 50000, "account_id": bank["id"]})
    loan = post(client, h, "/api/loans", {"name": "HDFC Term Loan", "lender": "HDFC", "interest_rate": 10, "opening_balance": 20000})
    post(client, h, f"/api/loans/{loan['id']}/txns", {"date": "2026-04-02", "type": "DISBURSEMENT", "principal": 100000, "account_id": bank["id"]})
    post(client, h, f"/api/loans/{loan['id']}/txns", {"date": "2026-05-02", "type": "EMI", "principal": 5000, "interest": 1000, "account_id": bank["id"]})
    post(client, h, f"/api/loans/{loan['id']}/txns", {"date": "2026-05-03", "type": "CHARGES", "interest": 500, "account_id": bank["id"]})

    cust = post(client, h, "/api/parties", {"name": "Local Retail", "gst_type": "REGISTERED", "gstin": gstin("27", "AAACL1234C"), "opening_balance": 3000})
    out_cust = post(client, h, "/api/parties", {"name": "Karnataka Buyer", "gst_type": "REGISTERED", "gstin": gstin("29", "AAACK5678D")})
    sup = post(client, h, "/api/parties", {"name": "Supplier Co", "type": "SUPPLIER", "gst_type": "REGISTERED",
                                           "gstin": gstin("27", "AAACW9999E"), "opening_balance": -2000})
    landlord = post(client, h, "/api/parties", {"name": "Landlord", "type": "SUPPLIER", "gst_type": "REGISTERED",
                                                "gstin": gstin("27", "AAACR4444R")})
    item = post(client, h, "/api/items", {"name": "Bottle", "hsn_sac": "7323", "unit": "PCS", "sale_price": 500,
                                          "purchase_price": 300, "gst_rate": 18, "opening_stock": 10,
                                          "opening_stock_date": "2026-04-01", "track_batch": True, "category": "Kitchen"})
    service = post(client, h, "/api/items", {"type": "SERVICE", "name": "Installation", "hsn_sac": "998719", "sale_price": 1000, "gst_rate": 18})
    line = lambda **k: {"item_id": item["id"], "name": "Bottle", "qty": 1, "rate": 500, "gst_rate": 18, **k}  # noqa: E731

    # purchase with TCS and batch
    post(client, h, "/api/vouchers", {"type": "PURCHASE", "date": "2026-04-05", "party_id": sup["id"], "tcs_rate": 0.1,
                                      "lines": [line(qty=20, rate=300, batch_no="B1", expiry_date="2027-12-31")]})
    # sale with TCS, serials
    s1 = post(client, h, "/api/vouchers", {"type": "SALE", "date": "2026-04-10", "party_id": cust["id"], "tcs_rate": 1,
                                           "lines": [line(qty=4, batch_no="B1", serial_nos="SN1,SN2,SN3,SN4", discount_pct=10),
                                                     {"item_id": service["id"], "name": "Installation", "qty": 1, "rate": 1000, "gst_rate": 18}]})
    assert s1["tcs_amount"] > 0
    # inter-state walk-in style fully paid via bank
    post(client, h, "/api/vouchers", {"type": "SALE", "date": "2026-04-12", "party_id": out_cust["id"],
                                      "lines": [line(qty=2)], "fully_paid": True, "payment_mode": "BANK",
                                      "payment_account_id": bank["id"]})
    # cash sale with no party
    post(client, h, "/api/vouchers", {"type": "SALE", "date": "2026-04-13", "lines": [line(qty=1)], "fully_paid": True})
    # credit note
    post(client, h, "/api/vouchers", {"type": "SALE_RETURN", "date": "2026-04-15", "party_id": cust["id"],
                                      "original_voucher_id": s1["id"], "lines": [line(qty=1, serial_nos="SN4")]})
    # debit note
    post(client, h, "/api/vouchers", {"type": "PURCHASE_RETURN", "date": "2026-04-16", "party_id": sup["id"],
                                      "lines": [line(qty=2, rate=300)]})

    # delivery challan -> invoice ; challan does not move stock
    stock_before = client.get(f"/api/items/{item['id']}", headers=h).json()["stock"]
    dc = post(client, h, "/api/vouchers", {"type": "DELIVERY_CHALLAN", "date": "2026-04-18", "party_id": cust["id"], "lines": [line(qty=3)]})
    assert dc["number"].startswith("DC/") and dc["status"] == "OPEN"
    assert client.get(f"/api/items/{item['id']}", headers=h).json()["stock"] == stock_before
    inv = post(client, h, "/api/vouchers", {"type": "SALE", "date": "2026-04-19", "party_id": cust["id"],
                                            "source_voucher_id": dc["id"], "lines": [line(qty=3)]})
    assert client.get(f"/api/vouchers/{dc['id']}", headers=h).json()["status"] == "CONVERTED"
    assert inv["source_voucher_id"] == dc["id"]
    # cannot convert twice
    r = client.post("/api/vouchers", headers=h, json={"type": "SALE", "date": "2026-04-19", "party_id": cust["id"],
                                                      "source_voucher_id": dc["id"], "lines": [line()]})
    assert r.status_code == 400

    so = post(client, h, "/api/vouchers", {"type": "SALE_ORDER", "date": "2026-04-20", "party_id": out_cust["id"], "lines": [line(qty=5)]})
    po = post(client, h, "/api/vouchers", {"type": "PURCHASE_ORDER", "date": "2026-04-20", "party_id": sup["id"], "lines": [line(qty=50, rate=290)]})
    post(client, h, "/api/vouchers", {"type": "PURCHASE", "date": "2026-04-22", "party_id": sup["id"],
                                      "source_voucher_id": po["id"], "lines": [line(qty=50, rate=290)]})
    # sale orders cannot take payments
    r = client.post("/api/vouchers", headers=h, json={"type": "SALE_ORDER", "date": "2026-04-20", "party_id": out_cust["id"],
                                                      "lines": [line()], "amount_paid": 10})
    assert r.status_code == 400

    # expenses
    cats = {c["name"]: c for c in client.get("/api/expenses/categories", headers=h).json()}
    assert "Rent" in cats and cats["Manufacturing Expenses"]["kind"] == "DIRECT"
    rent_item = post(client, h, "/api/expenses/items", {"name": "Shop rent", "category_id": cats["Rent"]["id"], "gst_rate": 18})
    tea = post(client, h, "/api/expenses/items", {"name": "Tea", "category_id": cats["Tea & Refreshments"]["id"]})
    rent = post(client, h, "/api/vouchers", {"type": "EXPENSE", "date": "2026-04-30", "party_id": landlord["id"],
                                             "expense_category_id": cats["Rent"]["id"], "tax_applicable": True,
                                             "lines": [{"expense_item_id": rent_item["id"], "name": "Shop rent", "qty": 1, "rate": 20000, "gst_rate": 18}]})
    assert rent["title"] == "Expense" and rent["cgst"] == 1800 and rent["status"] == "UNPAID"
    post(client, h, "/api/vouchers", {"type": "EXPENSE", "date": "2026-04-30", "expense_category_id": cats["Tea & Refreshments"]["id"],
                                      "lines": [{"expense_item_id": tea["id"], "name": "Tea", "qty": 30, "rate": 10}], "fully_paid": True})
    post(client, h, "/api/vouchers", {"type": "EXPENSE", "date": "2026-04-30", "expense_category_id": cats["Manufacturing Expenses"]["id"],
                                      "lines": [{"name": "Job work", "qty": 1, "rate": 1500}], "fully_paid": True, "payment_account_id": bank["id"]})
    # expense without party must be fully paid
    r = client.post("/api/vouchers", headers=h, json={"type": "EXPENSE", "date": "2026-04-30", "expense_category_id": cats["Rent"]["id"],
                                                      "lines": [{"name": "x", "qty": 1, "rate": 10}]})
    assert r.status_code == 400
    # pay rent with TDS (194-I 10%)
    pr = post(client, h, "/api/payments", {"type": "OUT", "date": "2026-05-01", "party_id": landlord["id"], "amount": 21600,
                                           "tds_amount": 2000, "account_id": bank["id"]})
    assert pr["allocated"] == 23600
    assert client.get(f"/api/vouchers/{rent['id']}", headers=h).json()["status"] == "PAID"

    # cheque from customer with TDS: party settled now, bank only when cleared
    chq = post(client, h, "/api/payments", {"type": "IN", "date": "2026-05-05", "party_id": cust["id"], "amount": 1000,
                                            "tds_amount": 50, "mode": "CHEQUE", "account_id": bank["id"], "reference": "000123"})
    assert chq["cheque_status"] == "OPEN"
    bank_bal = lambda: next(a for a in client.get("/api/accounts", headers=h).json() if a["id"] == bank["id"])["balance"]  # noqa: E731
    before = bank_bal()
    post(client, h, f"/api/cheques/{chq['id']}", {"action": "clear", "date": "2026-05-07"}, 200)
    assert bank_bal() == before + 1000
    # a bounced cheque restores the party's dues
    bal_before = client.get(f"/api/parties/{cust['id']}", headers=h).json()["balance"]
    bad = post(client, h, "/api/payments", {"type": "IN", "date": "2026-05-06", "party_id": cust["id"], "amount": 700, "mode": "CHEQUE"})
    assert client.get(f"/api/parties/{cust['id']}", headers=h).json()["balance"] == bal_before - 700
    post(client, h, f"/api/cheques/{bad['id']}", {"action": "bounce"}, 200)
    assert client.get(f"/api/parties/{cust['id']}", headers=h).json()["balance"] == bal_before
    # one cheque left open at period end
    post(client, h, "/api/payments", {"type": "IN", "date": "2026-05-08", "party_id": out_cust["id"], "amount": 300, "mode": "CHEQUE"})

    post(client, h, "/api/transfers", {"date": "2026-05-09", "from_account_id": bank["id"], "to_account_id": cash["id"], "amount": 2000})
    post(client, h, "/api/capital", {"date": "2026-05-10", "type": "DRAWINGS", "amount": 1000})
    post(client, h, "/api/tax-payments", {"date": "2026-05-20", "type": "GST", "amount": 100, "account_id": bank["id"]})
    post(client, h, "/api/tax-payments", {"date": "2026-05-20", "type": "TDS", "amount": 2000, "account_id": bank["id"]})
    post(client, h, f"/api/items/{item['id']}/adjust", {"qty": 1, "direction": "REDUCE", "date": "2026-05-25", "note": "broken"}, 200)
    return h, dict(cust=cust, sup=sup, item=item, bank=bank, loan=loan, cash=cash, cats=cats, so=so)


def run(client, h, slug, **params):
    r = client.get(f"/api/reports/run/{slug}", headers=h, params=params)
    assert r.status_code == 200, (slug, r.text)
    return r.json()


def test_balance_sheet_balances(client):
    h, _ = build_books(client)
    for as_of in ("2026-04-01", "2026-04-15", "2026-05-06", "2026-06-30"):
        bs = run(client, h, "balance-sheet", as_of=as_of)
        rows = [x for s in bs["sections"] for x in s["rows"]]
        assert not any("Difference" in x["label"] for x in rows), (as_of, rows)


def test_all_reports_run(client):
    h, ctx = build_books(client)
    catalog = client.get("/api/reports/catalog", headers=h).json()
    assert len(catalog) >= 50
    params = dict(date_from="2026-04-01", date_to="2026-06-30", as_of="2026-06-30", party_id=ctx["cust"]["id"],
                  item_id=ctx["item"]["id"], account_id=ctx["bank"]["id"], loan_id=ctx["loan"]["id"])
    for rep in catalog:
        if rep["href"]:
            continue
        out = run(client, h, rep["slug"], **params)
        assert out["title"] and isinstance(out["sections"], list), rep["slug"]


def test_report_numbers(client):
    h, ctx = build_books(client)
    p = {"date_from": "2026-04-01", "date_to": "2026-06-30"}
    pl = {x["label"].strip(): x["amount"] for x in run(client, h, "profit-loss", **p)["sections"][0]["rows"]}
    assert pl["Rent"] == 20000 and pl["Tea & Refreshments"] == 300
    assert pl["Direct expense: Manufacturing Expenses"] == 1500
    assert pl["Loan interest & charges"] == 1500
    loan = run(client, h, "loan-statement", loan_id=ctx["loan"]["id"], **p)
    assert loan["summary"][0]["value"] == 115000  # 20000 + 100000 - 5000
    tds = run(client, h, "tds-payable", **p)
    assert tds["summary"][0]["value"] == 2000
    assert run(client, h, "tds-receivable", **p)["summary"][0]["value"] == 50
    eq = run(client, h, "form-27eq", **p)
    assert eq["summary"][0]["value"] > 0
    batch = run(client, h, "batch", as_of="2026-06-30")["sections"][0]["rows"]
    assert batch[0]["batch"] == "B1" and batch[0]["balance"] == 16
    serial = run(client, h, "serial", as_of="2026-06-30")["sections"][0]["rows"]
    assert {s["serial"]: s["status"] for s in serial}["SN4"] == "In stock"
    orders = run(client, h, "sale-orders", **p)
    assert orders["summary"][1]["value"] == 1  # one open sale order
    cf = run(client, h, "cash-flow", **p)
    labels = {x["label"].strip(): x["amount"] for x in cf["sections"][0]["rows"]}
    assert labels["Loans received"] == 100000 and labels["Drawings"] == -1000
    dash = client.get("/api/reports/dashboard", headers=h).json()
    assert dash["inventory"]["items"] == 1 and "cash_bank" in dash and "expenses" in dash
    assert dash["cash_bank"]["cheques_in"] == 300


def test_import_templates_and_sales_import(client):
    h = make_business(client, signup(client))
    types = client.get("/api/import/types", headers=h).json()
    for t in types:
        r = client.get(f"/api/import/{t['entity']}/template", headers=h)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert "Data" in wb.sheetnames and "Instructions" in wb.sheetnames

    def csv_bytes(rows):
        buf = io.StringIO()
        csv.writer(buf).writerows(rows)
        return buf.getvalue().encode()

    head = ["Doc No", "Date", "Party Name", "Party GSTIN", "Item Name", "HSN/SAC", "Qty", "Unit", "Rate", "GST %", "Amount Paid"]
    good = csv_bytes([head,
                      ["IMP-1", "01/05/2026", "New Customer", gstin("29", "AAACN1111N"), "Imported Widget", "8471", "2", "PCS", "1000", "18", ""],
                      ["IMP-1", "01/05/2026", "", "", "Second Item", "8472", "1", "NOS", "500", "5", ""],
                      ["IMP-2", "02/05/2026", "", "", "Imported Widget", "", "1", "", "1000", "18", ""]])
    dry = client.post("/api/import/sales", headers=h, files={"file": ("s.csv", good)}, data={"dry_run": "true"}).json()
    assert dry["ok"] and not dry["saved"], dry
    assert client.get("/api/vouchers?type=SALE", headers=h).json() == []
    real = client.post("/api/import/sales", headers=h, files={"file": ("s.csv", good)}, data={"dry_run": "false"}).json()
    assert real["saved"] and real["summary"] == "2 documents", real
    vs = {v["number"]: v for v in client.get("/api/vouchers?type=SALE", headers=h).json()}
    assert vs["IMP-1"]["inter_state"] and vs["IMP-1"]["igst"] == 385
    assert vs["IMP-2"]["status"] == "PAID"  # walk-in sale auto fully paid

    bad = csv_bytes([head, ["IMP-3", "31/31/2026", "X", "", "Y", "", "1", "", "abc", "18", ""]])
    res = client.post("/api/import/sales", headers=h, files={"file": ("b.csv", bad)}, data={"dry_run": "false"}).json()
    assert not res["ok"] and res["errors"] and not res["saved"]
    assert len(client.get("/api/vouchers?type=SALE", headers=h).json()) == 2

    # items import updates existing and creates new
    items_csv = csv_bytes([["Item Name", "Item Code", "HSN/SAC", "Sale Price", "GST %", "Opening Stock"],
                           ["Imported Widget", "W1", "8471", "1200", "18", ""],
                           ["Brand New", "BN1", "3923", "50", "5", "100"]])
    res = client.post("/api/import/items", headers=h, files={"file": ("i.csv", items_csv)}, data={"dry_run": "false"}).json()
    assert res["saved"] and res["summary"] == "1 items created, 1 updated", res
    items = {i["name"]: i for i in client.get("/api/items", headers=h).json()}
    assert items["Imported Widget"]["sale_price"] == 1200 and items["Brand New"]["stock"] == 100


def test_hsn_master_and_tax_slab(client):
    h = make_business(client, signup(client))
    it = post(client, h, "/api/items", {"name": "Soap", "hsn_sac": "3401", "gst_rate": 18})
    hsn_csv = b"HSN/SAC Code,Description,GST %\n3401,Soap,5\n9983,Other professional services,18\n"
    res = client.post("/api/import/hsn", headers=h, files={"file": ("h.csv", hsn_csv)}, data={"dry_run": "false"}).json()
    assert res["saved"], res
    applied = post(client, h, "/api/hsn/apply-to-items", {}, 200)
    assert applied["updated"] == 1 and client.get(f"/api/items/{it['id']}", headers=h).json()["gst_rate"] == 5
    prev = post(client, h, "/api/tax-slab/update", {"from_rate": 5, "new_rate": 12, "preview": True}, 200)
    assert prev["count"] == 1 and not prev["applied"]
    post(client, h, "/api/tax-slab/update", {"from_rate": 5, "new_rate": 12, "preview": False}, 200)
    assert client.get(f"/api/items/{it['id']}", headers=h).json()["gst_rate"] == 12


def test_backup_restore_roundtrip(client):
    h, _ = build_books(client)
    b = post(client, h, "/api/backups", {})
    blob = client.get(f"/api/backups/{b['id']}/download", headers=h).content
    auth = {"Authorization": h["Authorization"]}
    restored = client.post("/api/backups/restore", headers=auth, files={"file": ("x.gstbak", blob)}).json()
    h2 = {**auth, "X-Business-Id": restored["id"]}
    assert "(restored)" in restored["name"]
    orig_bs = run(client, h, "balance-sheet", as_of="2026-06-30")
    new_bs = run(client, h2, "balance-sheet", as_of="2026-06-30")
    assert orig_bs["summary"] == new_bs["summary"]
    assert len(client.get("/api/vouchers", headers=h2).json()) == len(client.get("/api/vouchers", headers=h).json())
    # restored company is independent: cancelling there leaves the original untouched
    v = client.get("/api/vouchers?type=SALE", headers=h2).json()[0]
    post(client, h2, f"/api/vouchers/{v['id']}/cancel", {}, 200)
    assert all(not x["cancelled"] for x in client.get("/api/vouchers?type=SALE", headers=h).json())
    assert client.post("/api/backups/restore", headers=auth, files={"file": ("x.gstbak", b"junk")}).status_code == 400


def test_members(client):
    owner = make_business(client, signup(client, "owner@x.in"))
    staff = signup(client, "staff@x.in")
    post(client, owner, "/api/members", {"email": "staff@x.in", "role": "ACCOUNTANT"})
    h = {**staff, "X-Business-Id": owner["X-Business-Id"]}
    assert client.get("/api/parties", headers=h).status_code == 200
    assert client.post("/api/parties", headers=h, json={"name": "x"}).status_code == 403  # accountant is read-only
    assert client.post("/api/members", headers=owner, json={"email": "nobody@x.in", "role": "BILLING"}).status_code == 404
