"""Platform HSN/SAC master: official Excel / PDF uploads, template, bulk rates, requests, strict mode."""

import io

from openpyxl import Workbook, load_workbook

from app.services.hsn_master import clean_code
from tests.test_api_flow import make_business, signup
from tests.test_modules import post
from tests.test_security_admin import superadmin

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def portal_xlsx() -> bytes:
    """Same layout as the GST portal download: sheets HSN_MSTR and SAC_MSTR."""
    wb = Workbook()
    ws = wb.active
    ws.title = "HSN_MSTR"
    ws.append(["HSN_CD", "HSN_Description"])
    ws.append([1, "LIVE ANIMALS"])                    # number cell, leading zero lost → 01
    ws.append([101, "Live horses, asses, mules"])     # → 0101
    ws.append(["61", "Apparel, knitted or crocheted"])
    ws.append(["6109", "T-shirts, singlets and other vests"])
    ws.append(["61091000", "Of cotton"])
    ws.append(["ABC", "junk"])
    sac = wb.create_sheet("SAC_MSTR")
    sac.append(["SAC_CD", "SAC_Description"])
    sac.append(["99", "All Services"])
    sac.append(["9954", "Construction services"])
    sac.append(["99541", "Construction services of buildings"])
    sac.append(["995411", "Construction services of single dwelling or multi dwelling"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def text_pdf(lines: list[str]) -> bytes:
    """A minimal one-page PDF with one text line per row."""
    ops = ["BT", "/F1 9 Tf", "11 TL", "40 800 Td"]
    for line in lines:
        ops.append("(" + line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") Tj T*")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def upload(client, root, name, content, dry_run=True):
    r = client.post("/api/admin/hsn/import", headers=root, data={"dry_run": str(dry_run).lower()},
                    files={"file": (name, content, "application/octet-stream")})
    assert r.status_code == 200, r.text
    return r.json()


def test_clean_code():
    assert clean_code("0101 21 00") == "01012100" and clean_code("0101.21.00") == "01012100"
    assert clean_code(101, numeric_cell=True) == "0101" and clean_code(99541, numeric_cell=True) == "99541"
    assert clean_code("1") is None and clean_code("123456789") is None and clean_code("abc") is None


def test_portal_excel_pdf_template_and_rates(client, monkeypatch):
    root = superadmin(client, monkeypatch)

    t = client.get("/api/admin/hsn/template", headers=root)
    assert t.status_code == 200
    wb = load_workbook(io.BytesIO(t.content))
    assert wb.sheetnames == ["HSN SAC master", "Instructions"]
    assert [c.value for c in wb.active[1]] == ["HSN/SAC Code", "Description", "GST %", "Cess %", "Effective From"]
    # the template itself imports cleanly
    assert upload(client, root, "template.xlsx", t.content)["error_count"] == 0

    prev = upload(client, root, "HSN_SAC.xlsx", portal_xlsx())
    assert prev["dry_run"] and prev["format"] == "Excel" and prev["rows"] == 9 and prev["error_count"] == 1
    assert prev["hsn"] == 5 and prev["sac"] == 4 and prev["created"] == 9
    assert client.get("/api/admin/hsn", headers=root).json()["total"] == 0  # preview saved nothing

    done = upload(client, root, "HSN_SAC.xlsx", portal_xlsx(), dry_run=False)
    assert done["created"] == 9 and done["stats"]["without_rate"] == 9
    codes = {r["code"] for r in client.get("/api/admin/hsn", headers=root, params={"limit": 50}).json()["rows"]}
    assert {"01", "0101", "61091000", "99541"} <= codes

    # bulk rate for a chapter (codes of 4+ digits only), then a specific override
    assert post(client, root, "/api/admin/hsn/bulk-rate", {"prefix": "61", "gst_rate": 5}, 200)["updated"] == 2
    assert post(client, root, "/api/admin/hsn/bulk-rate", {"prefix": "61", "gst_rate": 12}, 200)["updated"] == 0  # only missing
    assert post(client, root, "/api/admin/hsn/bulk-rate", {"prefix": "9954", "gst_rate": 18}, 200)["updated"] == 3
    r = client.post("/api/admin/hsn/bulk-rate", headers=root, json={"prefix": "61", "gst_rate": 13})
    assert r.status_code == 422
    listing = client.get("/api/admin/hsn", headers=root, params={"missing_rate": True}).json()
    assert listing["matches"] == 4 and {x["code"] for x in listing["rows"]} == {"01", "0101", "61", "99"}
    assert client.get("/api/admin/hsn", headers=root, params={"kind": "SAC"}).json()["matches"] == 4

    # re-upload without rates keeps the rates already set
    again = upload(client, root, "HSN_SAC.xlsx", portal_xlsx(), dry_run=False)
    assert again["created"] == 0 and again["rate_changes"] == 0
    assert client.get("/api/admin/hsn", headers=root, params={"search": "6109"}).json()["rows"][0]["gst_rate"] == 5

    # PDF (e.g. the SAC scheme / tariff PDFs): codes + descriptions, never rates
    pdf = text_pdf([
        "Scheme of Classification of Services",
        "S.No. Chapter, Section, Heading or Group Service Code (Tariff) Service Description",
        "1 Heading 9965 Goods Transport Services",
        "2 Group 99651 Land transport services of goods",
        "3 Service Code 996511 Road transport services of Goods including letters, parcels, live animals,",
        "household and office furniture, containers etc. by refrigerator vehicles",
        "0402 10 10 -- Milk powder in packing of 25 kg or more 30%",
        "Page 2 of 40",
    ])
    p = upload(client, root, "sac.pdf", pdf)
    assert p["format"] == "PDF" and p["rows"] == 4 and p["with_rate"] == 0
    by = {r["code"]: r for r in p["sample"]}
    assert by["996511"]["description"].endswith("by refrigerator vehicles")
    assert by["04021010"]["gst_rate"] is None and "Milk powder" in by["04021010"]["description"]

    exp = client.get("/api/admin/hsn/export", headers=root)
    assert load_workbook(io.BytesIO(exp.content)).active.max_row == 10  # header + 9 codes


def test_requests_strict_mode_and_checks(client, monkeypatch):
    root = superadmin(client, monkeypatch)
    h = make_business(client, signup(client))
    upload(client, root, "HSN_SAC.xlsx", portal_xlsx(), dry_run=False)
    post(client, root, "/api/admin/hsn/bulk-rate", {"prefix": "61", "gst_rate": 5}, 200)

    look = client.get("/api/hsn-master/lookup/61091000", headers=h).json()
    assert look["found"] and look["gst_rate"] == 5 and [p["code"] for p in look["parents"]] == ["61", "6109"]
    assert client.get("/api/hsn-master/lookup/99541100", headers=h).json()["found"] is False

    # goods / service consistency is always checked
    r = client.post("/api/items", headers=h, json={"type": "SERVICE", "name": "Repair", "hsn_sac": "6109", "gst_rate": 18})
    assert r.status_code == 422 and "SAC" in r.text
    r = client.post("/api/items", headers=h, json={"type": "GOODS", "name": "Shirt", "hsn_sac": "995411", "gst_rate": 18})
    assert r.status_code == 422

    # strict off: unknown codes allowed
    post(client, h, "/api/items", {"name": "Widget", "hsn_sac": "84713010", "gst_rate": 18})
    post(client, root, "/api/admin/config/versions", {"effective_from": "2017-07-01", "values": {"hsn_strict": True}})
    r = client.post("/api/items", headers=h, json={"name": "Gadget", "hsn_sac": "85171300", "gst_rate": 18})
    assert r.status_code == 422 and "not in the official master" in r.text
    post(client, h, "/api/items", {"name": "Tee", "hsn_sac": "61091000", "gst_rate": 5})
    r = client.post("/api/vouchers", headers=h, json={"type": "SALE", "date": "2026-09-10", "fully_paid": True,
                    "lines": [{"name": "Gadget", "hsn_sac": "85171300", "qty": 1, "rate": 100, "gst_rate": 18}]})
    assert r.status_code == 400 and "official master" in r.text

    # request the missing code → super admin approves → allowed for everyone
    req = post(client, h, "/api/hsn-master/requests", {"code": "85171300", "description": "Smartphones", "gst_rate": 18})
    assert client.post("/api/hsn-master/requests", headers=h, json={"code": "85171300", "description": "Smartphones"}).status_code == 409
    assert client.post("/api/hsn-master/requests", headers=h, json={"code": "6109", "description": "dup"}).status_code == 409
    pending = client.get("/api/admin/hsn-requests", headers=root).json()
    assert len(pending) == 1 and pending[0]["business"] == "Sharma Traders"
    assert client.get("/api/admin/hsn", headers=root).json()["pending_requests"] == 1
    post(client, root, f"/api/admin/hsn-requests/{req['id']}", {"approve": True, "description": "Smartphones (mobile phones)"}, 200)
    assert client.get("/api/hsn-master/requests", headers=h).json()[0]["status"] == "APPROVED"
    look = client.get("/api/hsn-master/lookup/85171300", headers=h).json()
    assert look["found"] and look["gst_rate"] == 18 and look["description"] == "Smartphones (mobile phones)"
    post(client, h, "/api/items", {"name": "Gadget", "hsn_sac": "85171300", "gst_rate": 18})

    # rejecting
    req2 = post(client, h, "/api/hsn-master/requests", {"code": "12345678", "description": "Made up"})
    post(client, root, f"/api/admin/hsn-requests/{req2['id']}", {"approve": False, "admin_note": "Not a valid code"}, 200)
    mine = {r["code"]: r for r in client.get("/api/hsn-master/requests", headers=h).json()}
    assert mine["12345678"]["status"] == "REJECTED" and mine["12345678"]["admin_note"] == "Not a valid code"
    assert client.post(f"/api/admin/hsn-requests/{req2['id']}", headers=root, json={"approve": True}).status_code == 400

    # copying into the business list skips codes without a rate
    res = post(client, h, "/api/hsn-master/copy", {"codes": ["61091000", "0101", "7777"]}, 200)
    assert res == {"copied": 1, "without_rate": ["0101"], "not_found": ["7777"]}
