"""Excel / PDF export of any table."""

import io

from openpyxl import load_workbook

from app.services.table_export import indian
from tests.test_api_flow import make_business, signup

DOC = {
    "title": "Sale report", "subtitle": "01-09-2026 to 30-09-2026", "filename": "sale-report",
    "summary": [{"label": "Total", "value": 1234567.5, "type": "money"}],
    "sections": [{
        "title": "Invoices",
        "columns": [{"key": "date", "label": "Date", "type": "date"}, {"key": "party", "label": "Party", "type": "text"},
                    {"key": "amount", "label": "Amount", "type": "money"}, {"key": "qty", "label": "Qty", "type": "qty"}],
        "rows": [{"date": "2026-09-10", "party": "Karnataka Retail & Co <b>", "amount": 1180, "qty": 2},
                 {"date": "2026-09-11", "party": "हिंदी ग्राहक", "amount": -50.25, "qty": 1.5, "_style": "bold"}],
        "total": {"amount": 1129.75, "qty": 3.5}, "note": "Net of credit notes",
    }],
}


def test_indian_grouping():
    from decimal import Decimal
    assert indian(Decimal("1234567.5")) == "12,34,567.50" and indian(Decimal("-999")) == "-999.00"


def test_export_xlsx_and_pdf(client):
    h = make_business(client, signup(client))
    r = client.post("/api/export/table", headers=h, params={"format": "xlsx"}, json=DOC)
    assert r.status_code == 200 and 'filename="sale-report.xlsx"' in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    values = [[c for c in row if c is not None] for row in ws.iter_rows(values_only=True)]
    assert ["Sale report"] in values and any("Sharma Traders" in str(v) for v in values[1])
    assert ["Date", "Party", "Amount", "Qty"] in values
    assert any(row[:1] == ["Total"] and 1129.75 in row for row in values)

    r = client.post("/api/export/table", headers=h, params={"format": "pdf"}, json=DOC)
    assert r.status_code == 200 and r.content[:5] == b"%PDF-" and r.headers["content-type"] == "application/pdf"
    # works without a business too (admin pages), but needs a login
    assert client.post("/api/export/table", headers={"Authorization": h["Authorization"]}, params={"format": "pdf"}, json=DOC).status_code == 200
    assert client.post("/api/export/table", params={"format": "pdf"}, json=DOC).status_code == 401
