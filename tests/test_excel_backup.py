"""Backup as a zip of Excel files: readable files, exact restore when unchanged, rebuild from edited Excel files."""

import io
import json
import zipfile

from openpyxl import load_workbook

from tests.test_modules import build_books, post


def _zip(client, h):
    b = post(client, h, "/api/backups", {})
    r = client.get(f"/api/backups/{b['id']}/excel", headers=h)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/zip" and b["excel_filename"].endswith("-excel.zip")
    return r.content


def _restore(client, auth, blob, **form):
    return client.post("/api/backups/restore", headers=auth, files={"file": ("backup.zip", blob)},
                       data={k: str(v).lower() if isinstance(v, bool) else v for k, v in form.items()})


def _totals(client, h):
    vs = client.get("/api/vouchers", headers=h).json()
    by_type = {}
    for v in vs:
        if not v["cancelled"]:
            by_type[v["type"]] = round(by_type.get(v["type"], 0) + v["grand_total"], 2)
    parties = {p["name"]: p["balance"] for p in client.get("/api/parties", headers=h).json()}
    stock = {i["name"]: i.get("stock") for i in client.get("/api/items", headers=h).json()}
    return by_type, parties, stock


def test_excel_zip_readable_and_restores(client):
    h, _ = build_books(client)
    post(client, h, "/api/employees", {"name": "Ravi Kumar", "code": "E1", "salary_type": "MONTHLY", "salary": 20000, "pf": True})
    blob = _zip(client, h)
    z = zipfile.ZipFile(io.BytesIO(blob))
    names = z.namelist()
    assert "00 READ ME.txt" in names and "backup.gstbak" in names and "manifest.json" in names
    assert any("Business details" in n for n in names) and any("Sales invoices" in n for n in names)
    assert any("Staff" in n for n in names)
    sales = next(n for n in names if "Sales invoices" in n)
    ws = load_workbook(io.BytesIO(z.read(sales)))["Data"]
    heads = [c.value for c in ws[1]]
    assert heads[0] == "Doc No" and "Document Total  (for reading)" in heads
    assert ws.max_row > 2

    auth = {"Authorization": h["Authorization"]}
    # unchanged → the exact backup inside the zip
    r = _restore(client, auth, blob)
    assert r.status_code == 201, r.text
    assert r.json()["method"] == "exact"
    h_exact = {**auth, "X-Business-Id": r.json()["id"]}
    assert _totals(client, h_exact) == _totals(client, h)
    assert [e["name"] for e in client.get("/api/employees", headers=h_exact).json()] == ["Ravi Kumar"]  # staff now in backups

    # edit a party's phone in Excel → rebuilt from the Excel files
    parties_file = next(n for n in names if "Customers and suppliers" in n)
    wb = load_workbook(io.BytesIO(z.read(parties_file)))
    ws = wb["Data"]
    phone_col = [c.value for c in ws[1]].index("Phone") + 1
    ws.cell(row=2, column=phone_col, value="9811122233")
    buf = io.BytesIO()
    wb.save(buf)
    edited = io.BytesIO()
    with zipfile.ZipFile(edited, "w") as out:
        for n in names:
            out.writestr(n, buf.getvalue() if n == parties_file else z.read(n))
    check = _restore(client, auth, edited.getvalue(), dry_run=True)
    assert check.status_code == 201 and check.json()["ok"] and check.json()["method"] == "excel", check.text
    assert parties_file in check.json()["edited"]
    r = _restore(client, auth, edited.getvalue(), name="Rebuilt from Excel")
    assert r.status_code == 201, r.text
    h_xl = {**auth, "X-Business-Id": r.json()["id"]}
    orig_types, orig_parties, orig_stock = _totals(client, h)
    xl_types, xl_parties, xl_stock = _totals(client, h_xl)
    for t in ("SALE", "PURCHASE", "SALE_RETURN", "PURCHASE_RETURN"):
        assert xl_types.get(t) == orig_types.get(t), t
    assert xl_parties == orig_parties
    assert xl_stock == orig_stock
    assert any(p["phone"] == "9811122233" for p in client.get("/api/parties", headers=h_xl).json())
    assert [e["name"] for e in client.get("/api/employees", headers=h_xl).json()] == ["Ravi Kumar"]

    # a broken Excel file: nothing is saved, errors name the file and row
    before = len(client.get("/api/auth/me", headers=auth).json()["businesses"])
    sales_wb = load_workbook(io.BytesIO(z.read(sales)))
    sales_wb["Data"].cell(row=2, column=[c.value for c in sales_wb["Data"][1]].index("Qty") + 1, value="lots")
    buf2 = io.BytesIO()
    sales_wb.save(buf2)
    bad = io.BytesIO()
    with zipfile.ZipFile(bad, "w") as out:
        for n in names:
            out.writestr(n, buf2.getvalue() if n == sales else z.read(n))
    r = _restore(client, auth, bad.getvalue())
    assert r.status_code == 422 and r.json()["detail"]["errors"][0]["file"] == sales
    assert len(client.get("/api/auth/me", headers=auth).json()["businesses"]) == before

    # forcing exact restore on a zip without the backup inside
    noexact = io.BytesIO()
    with zipfile.ZipFile(noexact, "w") as out:
        for n in names:
            if n != "backup.gstbak":
                out.writestr(n, z.read(n))
    assert _restore(client, auth, noexact.getvalue(), mode="exact").status_code == 400
    manifest = json.loads(z.read("manifest.json"))
    assert manifest["format"] == "gst-billing-excel-backup"


def test_backup_carries_database_images(client, monkeypatch):
    """Logos kept in our database travel with the backup; the restored company points at its own copy."""
    from tests.test_image_storage import png
    from tests.test_phase2 import setup
    from tests.test_security_admin import superadmin

    root = superadmin(client, monkeypatch)
    client.put("/api/admin/documents-settings", headers=root, json={"images_storage": "DATABASE"})
    h, _, _ = setup(client)
    url = client.post("/api/uploads", headers=h, data={"kind": "logo"}, files={"file": ("l.png", png(), "image/png")}).json()["url"]
    biz = client.get("/api/businesses/current", headers=h).json()
    keep = {k: x for k, x in biz.items() if k not in ("id", "plan", "einvoice_password_set", "ewb_password_set", "gst_portal", "modules")}
    client.put("/api/businesses/current", headers=h, json={**keep, "logo_url": url})
    b = post(client, h, "/api/backups", {})
    blob = client.get(f"/api/backups/{b['id']}/download", headers=h).content
    auth = {"Authorization": h["Authorization"]}
    new = client.post("/api/backups/restore", headers=auth, files={"file": ("x.gstbak", blob)}).json()
    new_logo = client.get("/api/businesses/current", headers={**auth, "X-Business-Id": new["id"]}).json()["logo_url"]
    assert new_logo != url and new_logo.startswith("/api/files/img/")
    assert client.get(new_logo).content == client.get(url).content


def test_super_admin_switches_excel_backup(client, monkeypatch):
    from tests.test_phase2 import setup
    from tests.test_security_admin import superadmin

    root = superadmin(client, monkeypatch)
    h, _, _ = setup(client)
    bid = h["X-Business-Id"]
    b = post(client, h, "/api/backups", {})
    assert client.get("/api/backups/settings", headers=h).json() == {"excel": True}  # on by default

    # off for this business only
    r = client.put(f"/api/admin/businesses/{bid}/excel-backup", headers=root, json={"enabled": False})
    assert r.json() == {"excel_backup": False, "excel_backup_on": False}
    assert client.get("/api/backups/settings", headers=h).json() == {"excel": False}
    assert client.get(f"/api/backups/{b['id']}/excel", headers=h).status_code == 403
    assert client.get(f"/api/backups/{b['id']}/download", headers=h).status_code == 200  # normal backups unaffected
    zip_blob = b"PK\x03\x04" + b"x" * 20
    auth = {"Authorization": h["Authorization"]}
    assert client.post("/api/backups/restore", headers=auth, files={"file": ("b.zip", zip_blob)}).status_code == 403

    # platform default off, this business switched on
    assert client.put("/api/admin/excel-backup", headers=root, json={"enabled": False}).json() == {"default": False}
    client.put(f"/api/admin/businesses/{bid}/excel-backup", headers=root, json={"enabled": True})
    assert client.get(f"/api/backups/{b['id']}/excel", headers=h).status_code == 200
    # back to the platform default (off)
    client.put(f"/api/admin/businesses/{bid}/excel-backup", headers=root, json={"enabled": None})
    row = next(x for x in client.get("/api/admin/businesses", headers=root).json() if x["id"] == bid)
    assert row["excel_backup"] is None and row["excel_backup_on"] is False


def test_zip_of_extracted_folder(client):
    """People extract the zip, edit a file, and zip the whole folder again: files then sit inside a folder."""
    from tests.test_api_flow import make_business, signup

    h = make_business(client, signup(client, "folder@x.in"))
    post(client, h, "/api/expenses/items", {"name": "Tea", "category_id": client.get("/api/expenses/categories", headers=h).json()[0]["id"]})
    blob = _zip(client, h)
    z = zipfile.ZipFile(io.BytesIO(blob))
    biz_file = next(n for n in z.namelist() if "Business details" in n)
    wb = load_workbook(io.BytesIO(z.read(biz_file)))
    ws = wb["Data"]
    ws.cell(row=2, column=[c.value for c in ws[1]].index("Address") + 1, value="Shop 20, New Market Road")
    buf = io.BytesIO()
    wb.save(buf)
    auth = {"Authorization": h["Authorization"]}
    for keep_manifest in (True, False):  # without the manifest, files are recognised by their names
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as o:
            for n in z.namelist():
                if n == "manifest.json" and not keep_manifest:
                    continue
                o.writestr(f"My backup/{n}", buf.getvalue() if n == biz_file else z.read(n))
        r = _restore(client, auth, out.getvalue(), name=f"Folder {keep_manifest}")
        assert r.status_code == 201, r.text
        assert r.json()["method"] == "excel"
        h2 = {**auth, "X-Business-Id": r.json()["id"]}
        assert client.get("/api/businesses/current", headers=h2).json()["address"] == "Shop 20, New Market Road"
        assert [e["name"] for e in client.get("/api/expenses/items", headers=h2).json()].count("Tea") == 1
